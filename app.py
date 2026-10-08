import os, math, json, uuid, csv, io, re, statistics
from pathlib import Path
from functools import wraps
from datetime import datetime, timedelta
from openpyxl import load_workbook, Workbook
try:
    import xlrd
except ImportError:
    xlrd = None
from flask_wtf.csrf import CSRFProtect, CSRFError
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file, jsonify, g, make_response
import mysql.connector
from mysql.connector import Error
from mysql.connector.pooling import MySQLConnectionPool
from dotenv import load_dotenv
from werkzeug.security import generate_password_hash, check_password_hash
from large_import_config import MAX_IMPORT_ROWS, IMPORT_BATCH_SIZE, validate_import_row_count

APP_VERSION = "32.6"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, '.env'), override=True)
app = Flask(__name__)
csrf = CSRFProtect(app)
limiter = Limiter(key_func=get_remote_address, app=app, default_limits=[], storage_uri=os.getenv('RATELIMIT_STORAGE_URI','memory://'))
DB_POOL = None
DB_POOL_SIZE = max(3, min(20, int(os.getenv("DB_POOL_SIZE", "8"))))
SECRET_KEY = os.getenv("SECRET_KEY")
if not SECRET_KEY:
    raise RuntimeError("SECRET_KEY is not configured. Create a private .env file before starting the application.")
app.secret_key = SECRET_KEY
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE=os.getenv("SESSION_COOKIE_SAMESITE", "Lax"),
    SESSION_COOKIE_SECURE=os.getenv("SESSION_COOKIE_SECURE", "0").lower() in {"1","true","yes"},
    SESSION_REFRESH_EACH_REQUEST=False,
    WTF_CSRF_TIME_LIMIT=int(os.getenv("WTF_CSRF_TIME_LIMIT", "3600")),
    MAX_CONTENT_LENGTH=int(os.getenv("MAX_CONTENT_LENGTH", str(1024 * 1024 * 1024))),
    SEND_FILE_MAX_AGE_DEFAULT=0,
    JSON_SORT_KEYS=False,
)
# Large academic imports: support college-scale files while retaining a sane request ceiling.
IMPORT_DIR = os.path.join(BASE_DIR, 'instance', 'imports')
os.makedirs(IMPORT_DIR, exist_ok=True)

# ---------------- DB / SETTINGS ----------------
def get_db():
    global DB_POOL
    if DB_POOL is None:
        DB_POOL = MySQLConnectionPool(
            pool_name='ams_pool',
            pool_size=DB_POOL_SIZE,
            pool_reset_session=True,
            host=os.getenv('DB_HOST'),
            user=os.getenv('DB_USER'),
            password=os.getenv('DB_PASSWORD'),
            database=os.getenv('DB_NAME'),
            autocommit=False,
        )
    return DB_POOL.get_connection()

def query_db(sql, params=(), fetchone=False, commit=False):
    conn = cursor = None
    try:
        conn = get_db(); cursor = conn.cursor(dictionary=True)
        cursor.execute(sql, params)
        if commit:
            conn.commit(); return cursor.lastrowid
        return cursor.fetchone() if fetchone else cursor.fetchall()
    finally:
        if cursor: cursor.close()
        if conn: conn.close()

def ensure_column(cursor, table, column, definition):
    cursor.execute("SELECT COUNT(*) FROM information_schema.columns WHERE table_schema=DATABASE() AND table_name=%s AND column_name=%s", (table, column))
    if cursor.fetchone()[0] == 0:
        cursor.execute(f'ALTER TABLE {table} ADD COLUMN {column} {definition}')


def ensure_auth_extensions(cursor):
    ensure_column(cursor, 'users', 'password_hint', 'VARCHAR(255) NULL')
    ensure_column(cursor, 'users', 'login_failures', 'INT NOT NULL DEFAULT 0')
    ensure_column(cursor, 'users', 'last_failed_at', 'DATETIME NULL')
    cursor.execute("""CREATE TABLE IF NOT EXISTS password_reset_requests (
        request_id INT AUTO_INCREMENT PRIMARY KEY,
        user_id INT NOT NULL,
        requested_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        status ENUM('pending','completed','dismissed') NOT NULL DEFAULT 'pending',
        FOREIGN KEY(user_id) REFERENCES users(user_id) ON DELETE CASCADE
    )""")

def ensure_schema_extensions():
    conn = cursor = None
    try:
        conn = get_db(); cursor = conn.cursor()
        cursor.execute("CREATE TABLE IF NOT EXISTS settings (setting_key VARCHAR(80) PRIMARY KEY, setting_value VARCHAR(255) NOT NULL)")
        ensure_column(cursor, 'students', 'enrollment_no', 'VARCHAR(50) NULL UNIQUE AFTER roll_no')
        ensure_column(cursor, 'users', 'student_id', 'INT NULL')
        ensure_column(cursor, 'students', 'teacher_remark', 'VARCHAR(500) NULL')
        ensure_auth_extensions(cursor)
        # Only alter the role enum when a legacy installation actually needs it.
        cursor.execute("SELECT COLUMN_TYPE FROM information_schema.columns WHERE table_schema=DATABASE() AND table_name='users' AND column_name='role'")
        role_type = cursor.fetchone()
        if role_type and 'student' not in str(role_type[0]):
            cursor.execute("ALTER TABLE users MODIFY role ENUM('admin','teacher','staff','student') NOT NULL DEFAULT 'teacher'")
        cursor.execute("CREATE TABLE IF NOT EXISTS teacher_classes (user_id INT NOT NULL, class_name VARCHAR(50) NOT NULL, PRIMARY KEY(user_id,class_name), FOREIGN KEY(user_id) REFERENCES users(user_id) ON DELETE CASCADE)")
        cursor.execute("CREATE TABLE IF NOT EXISTS attendance_months (attendance_month_id INT AUTO_INCREMENT PRIMARY KEY, month_key CHAR(7) NOT NULL UNIQUE, month_label VARCHAR(40) NOT NULL)")
        cursor.execute("CREATE TABLE IF NOT EXISTS attendance (attendance_id INT AUTO_INCREMENT PRIMARY KEY, student_id INT NOT NULL, attendance_month_id INT NOT NULL, subject_id INT NULL, working_days DECIMAL(8,2) NOT NULL DEFAULT 0, present_days DECIMAL(8,2) NOT NULL DEFAULT 0, notes VARCHAR(255) NULL, UNIQUE KEY uq_attendance(student_id,attendance_month_id,subject_id), FOREIGN KEY(student_id) REFERENCES students(student_id) ON DELETE CASCADE, FOREIGN KEY(attendance_month_id) REFERENCES attendance_months(attendance_month_id) ON DELETE CASCADE, FOREIGN KEY(subject_id) REFERENCES subjects(subject_id) ON DELETE SET NULL)")
        defaults=[('institution_type','college'),('institution_name','Academic Marks Management System'),('use_enrollment_no','0'),('detention_threshold','75')]
        for k,v in defaults:
            cursor.execute('INSERT IGNORE INTO settings(setting_key,setting_value) VALUES(%s,%s)', (k,v))
        # Lightweight indexes used by result/analysis/attendance/import queries.
        for sql in (
            'CREATE INDEX idx_students_class ON students(student_class)',
            'CREATE INDEX idx_marks_sem_exam ON marks(semester_id,exam_id)',
            'CREATE INDEX idx_ss_student_subject ON student_subjects(student_id,subject_id)',
            'CREATE INDEX idx_attendance_student_month ON attendance(student_id,attendance_month_id)',
            'CREATE INDEX idx_reset_status_user ON password_reset_requests(user_id,status)',
        ):
            try: cursor.execute(sql)
            except Error as e:
                if getattr(e, 'errno', None) != 1061: raise
        cursor.execute("INSERT INTO settings(setting_key,setting_value) VALUES('schema_extensions_version','32.4') ON DUPLICATE KEY UPDATE setting_value='32.4'")
        conn.commit()
    except Error:
        if conn: conn.rollback()
        raise
    finally:
        if cursor: cursor.close()
        if conn: conn.close()

def get_setting(key, default=''):
    try:
        row=query_db('SELECT setting_value FROM settings WHERE setting_key=%s',(key,),fetchone=True)
        return row['setting_value'] if row else default
    except Exception:
        return default

def institution_config():
    if hasattr(g,'institution_config_cache'): return g.institution_config_cache
    g.institution_config_cache={'type':get_setting('institution_type','college'),'name':get_setting('institution_name','Academic Marks Management System'),
            'use_enrollment':get_setting('use_enrollment_no','0')=='1','detention_threshold':float(get_setting('detention_threshold','75') or 75)}
    return g.institution_config_cache

def admin_exists():
    try:
        return bool(query_db("SELECT user_id FROM users WHERE role='admin' LIMIT 1", fetchone=True))
    except Exception:
        return False

@app.before_request
def init():
    if not app.config.get('SCHEMA_READY'):
        try:
            marker = query_db("SELECT setting_value FROM settings WHERE setting_key='schema_extensions_version'", fetchone=True)
            if marker and marker.get('setting_value') == APP_VERSION:
                app.config['SCHEMA_READY']=True
            else:
                ensure_schema_extensions(); app.config['SCHEMA_READY']=True
        except Exception as exc:
            app.logger.exception('Schema initialization failed: %s',exc)

@app.after_request
def security_headers(response):
    response.headers.setdefault('X-Content-Type-Options','nosniff')
    response.headers.setdefault('X-Frame-Options','SAMEORIGIN')
    response.headers.setdefault('Referrer-Policy','same-origin')
    if request.endpoint in {'login','forgot_password','setup_administrator'}:
        response.headers['Cache-Control']='no-store, no-cache, must-revalidate, max-age=0'
        response.headers['Pragma']='no-cache'
    return response

def current_user():
    if hasattr(g,'current_user_cache'): return g.current_user_cache
    uid=session.get('user_id')
    g.current_user_cache = query_db('SELECT u.user_id,u.full_name,u.username,u.role,u.is_active,u.student_id,s.roll_no,s.enrollment_no,s.student_class FROM users u LEFT JOIN students s ON s.student_id=u.student_id WHERE u.user_id=%s',(uid,),fetchone=True) if uid else None
    return g.current_user_cache

def teacher_classes(uid):
    key=f'teacher_classes_{uid}'
    if hasattr(g,key): return getattr(g,key)
    value=[x['class_name'] for x in query_db('SELECT class_name FROM teacher_classes WHERE user_id=%s ORDER BY class_name',(uid,))]
    setattr(g,key,value); return value

def allowed_classes():
    if hasattr(g,'allowed_classes_cache'): return g.allowed_classes_cache
    u=current_user()
    if not u: value=[]
    elif u['role']=='admin': value=[x['student_class'] for x in query_db("SELECT DISTINCT student_class FROM students WHERE student_class<>'' ORDER BY student_class")]
    elif u['role']=='teacher': value=teacher_classes(u['user_id'])
    elif u['role']=='student': value=[u['student_class']] if u.get('student_class') else []
    elif u['role']=='staff': value=[x['student_class'] for x in query_db("SELECT DISTINCT student_class FROM students WHERE student_class<>'' ORDER BY student_class")]
    else: value=[]
    g.allowed_classes_cache=value; return value

def active_class():
    if hasattr(g,'active_class_cache'): return g.active_class_cache
    u=current_user()
    if not u or u['role']!='teacher': value=None
    else:
        classes=teacher_classes(u['user_id']); selected=session.get('active_class')
        value=selected if selected in classes else (classes[0] if classes else None)
    g.active_class_cache=value; return value

@app.context_processor
def globals():
    u=current_user()
    return {'logged_in':bool(session.get('user_id')),'current_user':u,'institution':institution_config(),'active_class':active_class(),'allowed_classes':allowed_classes()}

def login_required(f):
    @wraps(f)
    def w(*a,**kw):
        if not session.get('user_id'): return redirect(url_for('login'))
        u=current_user()
        if not u or not u['is_active']:
            session.clear(); flash('Your account is inactive.','danger'); return redirect(url_for('login'))
        return f(*a,**kw)
    return w

def role_required(*roles):
    def deco(f):
        @wraps(f)
        def w(*a,**kw):
            u=current_user()
            if not u or u['role'] not in roles:
                flash('You do not have permission to access this page.','danger'); return redirect(url_for('index'))
            return f(*a,**kw)
        return w
    return deco

def enforce_class(class_name):
    u=current_user()
    if not u: return False
    if u['role']=='admin': return True
    if u['role']=='teacher': return class_name in teacher_classes(u['user_id'])
    if u['role']=='student': return class_name == u.get('student_class')
    return False

def safe_float(v):
    try:
        n=float(v); return n if math.isfinite(n) else None
    except (TypeError,ValueError): return None

# ---------------- AUTH / USERS ----------------
@app.route('/setup/administrator',methods=['GET','POST'])
def setup_administrator():
    if admin_exists():
        return redirect(url_for('login', role='admin'))
    if request.method=='POST':
        name=request.form.get('full_name','').strip()
        username=request.form.get('username','').strip()
        password=request.form.get('password','')
        confirm=request.form.get('confirm_password','')
        recovery=request.form.get('recovery_code','')
        if not name or not username or len(password)<8 or password!=confirm or len(recovery)<12:
            flash('Enter all fields. Password must be at least 8 characters and the recovery code at least 12 characters.','danger')
        elif query_db('SELECT user_id FROM users WHERE username=%s',(username,),fetchone=True):
            flash('That username is already in use. Choose another one.','danger')
        else:
            query_db('INSERT INTO users(full_name,username,password_hash,role) VALUES(%s,%s,%s,%s)',(name,username,generate_password_hash(password),'admin'),commit=True)
            # Store only a hash of the recovery code in the settings table; the plaintext is supplied by the deployment environment for normal recovery.
            query_db('INSERT INTO settings(setting_key,setting_value) VALUES(%s,%s) ON DUPLICATE KEY UPDATE setting_value=VALUES(setting_value)',('admin_recovery_code_hash',generate_password_hash(recovery)),commit=True)
            flash('Administrator account created. Keep the recovery code somewhere secure.','success')
            return redirect(url_for('login',role='admin'))
    return render_template('setup_admin.html')

@app.route('/login',methods=['GET','POST'])
@limiter.limit('10 per minute', methods=['POST'])
def login():
    selected_role=request.args.get('role') or request.form.get('login_role') or ''
    if selected_role not in ('admin','teacher','staff','student'): selected_role=''
    if selected_role=='admin' and not admin_exists():
        return redirect(url_for('setup_administrator'))
    if request.method=='POST':
        username=request.form.get('username','').strip()
        password=request.form.get('password','')
        u=query_db('SELECT * FROM users WHERE username=%s',(username,),fetchone=True)
        if u and u['role']!=selected_role:
            flash('Invalid username or password.','danger')
            return render_template('login.html',selected_role=selected_role)
        if u:
            failures=int(u.get('login_failures') or 0); last=u.get('last_failed_at')
            if failures >= 5 and last and datetime.now() - last < timedelta(minutes=10):
                flash('Too many unsuccessful attempts. Please wait a few minutes or use password recovery.','warning')
                return render_template('login.html',selected_role=selected_role)
        if u and u['is_active'] and check_password_hash(u['password_hash'],password):
            query_db('UPDATE users SET login_failures=0,last_failed_at=NULL WHERE user_id=%s',(u['user_id'],),commit=True)
            session.clear(); session['user_id']=u['user_id']; session['role']=u['role']
            if u['role']=='teacher':
                classes=teacher_classes(u['user_id'])
                if classes: session['active_class']=classes[0]
            flash(f"Welcome, {u['full_name']}.",'success')
            return redirect(url_for('select_class') if u['role']=='teacher' else url_for('index'))
        if u:
            failures=int(u.get('login_failures') or 0)+1
            query_db('UPDATE users SET login_failures=%s,last_failed_at=NOW() WHERE user_id=%s',(failures,u['user_id']),commit=True)
            if failures>=5:
                flash('Too many unsuccessful attempts. Use the password recovery option or contact an authorized administrator.', 'warning')
            else:
                flash('Invalid username or password.','danger')
        else:
            flash('Invalid username or password.','danger')
    return render_template('login.html',selected_role=selected_role)

@app.route('/forgot-password',methods=['GET','POST'])
@limiter.limit('5 per minute', methods=['POST'])
def forgot_password():
    role=request.form.get('role',request.args.get('role','teacher'))
    if role not in ('teacher','student','staff','admin'): role='teacher'
    if request.method=='POST':
        username=request.form.get('username','').strip()
        u=query_db('SELECT user_id,role FROM users WHERE username=%s',(username,),fetchone=True)
        if not u or u['role']!=role:
            flash('We could not find an account for that login type.','danger')
        elif role=='admin':
            code=request.form.get('recovery_code','')
            expected=os.getenv('ADMIN_RECOVERY_CODE','').strip()
            stored=query_db('SELECT setting_value FROM settings WHERE setting_key=%s',('admin_recovery_code_hash',),fetchone=True)
            valid=bool(code and expected and code==expected)
            if not valid and code and stored:
                valid=check_password_hash(stored['setting_value'],code)
            if valid:
                newpw=request.form.get('new_password','')
                if len(newpw)<8: flash('New password must be at least 8 characters.','danger')
                else:
                    query_db('UPDATE users SET password_hash=%s,login_failures=0 WHERE user_id=%s',(generate_password_hash(newpw),u['user_id']),commit=True)
                    flash('Administrator password reset successfully.','success'); return redirect(url_for('login',role='admin'))
            else: flash('Administrator recovery code is incorrect.','danger')
        else:
            # Lock the account row so two simultaneous submissions cannot create two
            # pending requests for the same user.
            conn = cur = None
            try:
                conn=get_db(); cur=conn.cursor(dictionary=True)
                cur.execute("SELECT user_id FROM users WHERE user_id=%s FOR UPDATE", (u['user_id'],))
                locked_user=cur.fetchone()
                cur.execute(
                    "SELECT request_id FROM password_reset_requests WHERE user_id=%s AND status='pending' LIMIT 1 FOR UPDATE",
                    (u['user_id'],)
                )
                pending=cur.fetchone()
                if pending:
                    flash('A password reset request is already pending. Please wait for it to be handled before sending another request.', 'warning')
                elif not locked_user:
                    flash('We could not process that request. Please try again.', 'danger')
                else:
                    cur.execute('INSERT INTO password_reset_requests(user_id) VALUES(%s)',(u['user_id'],))
                    conn.commit()
                    flash('Reset request sent to the administrator. You can send another request after this one is handled or deleted.','success')
            except Error:
                if conn: conn.rollback()
                flash('We could not send the reset request right now. Please try again.', 'danger')
            finally:
                if cur: cur.close()
                if conn: conn.close()
    return render_template('forgot_password.html',role=role)

@app.route('/select-class',methods=['GET','POST'])
@login_required
@role_required('teacher')
def select_class():
    classes=teacher_classes(current_user()['user_id'])
    if request.method=='POST':
        c=request.form.get('class_name','')
        if c in classes: session['active_class']=c; flash(f'Active class changed to {c}.','success'); return redirect(url_for('index'))
        flash('Please select one of your assigned classes.','danger')
    return render_template('select_class.html',classes=classes,selected=active_class())

@app.post('/switch-class')
@login_required
@role_required('teacher')
def switch_class():
    c=request.form.get('class_name','')
    if c in teacher_classes(current_user()['user_id']): session['active_class']=c
    else: flash('You are not assigned to that class.','danger')
    return redirect(url_for('index'))

@app.post('/logout')
@login_required
def logout():
    session.clear(); return redirect(url_for('login'))

@app.route('/users',methods=['GET','POST'])
@login_required
@role_required('admin')
def users():
    students=query_db('SELECT student_id,roll_no,enrollment_no,student_name,student_class FROM students ORDER BY student_class,student_name')
    classes=query_db("SELECT DISTINCT student_class FROM students WHERE student_class<>'' ORDER BY student_class")
    if request.method=='POST':
        name=request.form.get('full_name','').strip(); username=request.form.get('username','').strip(); pw=request.form.get('password',''); role=request.form.get('role','teacher'); student_id=request.form.get('student_id',type=int); 
        selected_classes=[x for x in request.form.getlist('classes') if x]
        if not name or not username or len(pw)<8 or role not in ('admin','teacher','staff','student'):
            flash('Enter valid user details. Password must be at least 8 characters.','danger')
        elif role=='student' and not student_id:
            flash('Select the student account to link.','danger')
        elif role=='teacher' and not selected_classes:
            flash('Select at least one class for the teacher.','danger')
        else:
            conn=cur=None
            try:
                conn=get_db(); cur=conn.cursor()
                cur.execute('INSERT INTO users(full_name,username,password_hash,role,student_id) VALUES(%s,%s,%s,%s,%s)',(name,username,generate_password_hash(pw),role,student_id if role=='student' else None)); uid=cur.lastrowid
                for c in selected_classes: cur.execute('INSERT INTO teacher_classes(user_id,class_name) VALUES(%s,%s)',(uid,c))
                conn.commit(); flash('User created successfully.','success')
            except Error as e:
                if conn: conn.rollback()
                flash('Username already exists or the user could not be created.','danger')
            finally:
                if cur: cur.close()
                if conn: conn.close()
    rows=query_db('SELECT u.user_id,u.full_name,u.username,u.role,u.is_active,u.created_at,s.student_name,s.roll_no,s.student_class FROM users u LEFT JOIN students s ON s.student_id=u.student_id ORDER BY u.user_id DESC')
    reset_requests=query_db('SELECT pr.request_id,pr.requested_at,u.full_name,u.username,u.role FROM password_reset_requests pr JOIN users u ON u.user_id=pr.user_id WHERE pr.status="pending" ORDER BY pr.requested_at DESC')
    for r in rows: r['assigned_classes']=teacher_classes(r['user_id']) if r['role']=='teacher' else []
    return render_template('users.html',users=rows,students=students,classes=classes,reset_requests=reset_requests)

@app.post('/users/<int:user_id>/toggle')
@login_required
@role_required('admin')
def toggle_user(user_id):
    if user_id==session.get('user_id'): flash('You cannot deactivate your own account.','danger')
    else: query_db('UPDATE users SET is_active=NOT is_active WHERE user_id=%s',(user_id,),commit=True); flash('User status updated.','success')
    return redirect(url_for('users'))

@app.post('/users/<int:user_id>/delete')
@login_required
@role_required('admin')
def delete_user(user_id):
    if user_id==session.get('user_id'): flash('You cannot delete your own account.','danger'); return redirect(url_for('users'))
    try: query_db('DELETE FROM users WHERE user_id=%s',(user_id,),commit=True); flash('User deleted successfully.','success')
    except Error: flash('Unable to delete this user.','danger')
    return redirect(url_for('users'))


@app.post('/users/reset-request/<int:request_id>')
@login_required
def complete_reset_request(request_id):
    u=current_user()
    req=query_db('SELECT pr.request_id,u.user_id,u.role,u.full_name,s.student_class FROM password_reset_requests pr JOIN users u ON u.user_id=pr.user_id LEFT JOIN students s ON s.student_id=u.student_id WHERE pr.request_id=%s AND pr.status="pending"',(request_id,),fetchone=True)
    if not req:
        flash('Reset request not found.','danger'); return redirect(url_for('users') if u and u['role']=='admin' else url_for('index'))
    if u['role']!='admin' and not (u['role']=='teacher' and req['role']=='student' and enforce_class(req['student_class'])):
        flash('You do not have permission to reset this account.','danger'); return redirect(url_for('index'))
    pw=request.form.get('new_password','')
    if len(pw)<8:
        flash('Password must be at least 8 characters.','danger'); return redirect(url_for('users') if u['role']=='admin' else url_for('teacher_password_resets'))
    query_db('UPDATE users SET password_hash=%s,login_failures=0 WHERE user_id=%s',(generate_password_hash(pw),req['user_id']),commit=True)
    query_db('UPDATE password_reset_requests SET status="completed" WHERE request_id=%s',(request_id,),commit=True)
    flash(f"Password reset for {req['full_name']}.",'success')
    return redirect(url_for('users') if u['role']=='admin' else url_for('teacher_password_resets'))

@app.post('/users/reset-request/<int:request_id>/delete')
@login_required
@role_required('admin')
def delete_reset_request(request_id):
    req=query_db(
        "SELECT request_id FROM password_reset_requests WHERE request_id=%s AND status='pending'",
        (request_id,), fetchone=True
    )
    if not req:
        flash('Reset request not found or it has already been handled.', 'danger')
    else:
        # Keep an audit trail while making the account eligible to submit a new request.
        query_db(
            "UPDATE password_reset_requests SET status='dismissed' WHERE request_id=%s AND status='pending'",
            (request_id,), commit=True
        )
        flash('Password reset request deleted. The user can submit a new request now.', 'success')
    return redirect(url_for('users'))

@app.route('/password-resets')
@login_required
@role_required('teacher')
def teacher_password_resets():
    cs=teacher_classes(current_user()['user_id'])
    if not cs: return render_template('teacher_resets.html',reset_requests=[])
    ph=','.join(['%s']*len(cs))
    sql=("SELECT pr.request_id,pr.requested_at,u.full_name,u.username,u.role,s.student_class "
         "FROM password_reset_requests pr JOIN users u ON u.user_id=pr.user_id "
         "JOIN students s ON s.student_id=u.student_id "
         f"WHERE pr.status='pending' AND u.role='student' AND s.student_class IN ({ph}) "
         "ORDER BY pr.requested_at DESC")
    reqs=query_db(sql,cs)
    return render_template('teacher_resets.html',reset_requests=reqs)

# ---------------- DASHBOARD ----------------
def student_scope_condition(alias='s'):
    u=current_user()
    if u['role']=='admin': return '1=1',[]
    if u['role']=='teacher':
        cs=teacher_classes(u['user_id']);
        if not cs: return '1=0',[]
        return f"{alias}.student_class IN ({','.join(['%s']*len(cs))})",cs
    if u['role']=='staff': return '1=1',[]
    return f'{alias}.student_id=%s',[u['student_id']]

@app.route('/')
@login_required
def index():
    scope,sp=student_scope_condition()
    stats={'students':query_db(f'SELECT COUNT(*) n FROM students s WHERE {scope}',sp,fetchone=True)['n'],
           'subjects':query_db('SELECT COUNT(*) n FROM subjects',fetchone=True)['n'],
           'marks':query_db(f'SELECT COUNT(*) n FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id JOIN students s ON s.student_id=ss.student_id WHERE {scope}',sp,fetchone=True)['n'],
           'attendance':0}
    att=query_db(f'''SELECT ROUND(SUM(a.present_days)/NULLIF(SUM(a.working_days),0)*100,2) pct FROM attendance a JOIN students s ON s.student_id=a.student_id WHERE {scope}''',sp,fetchone=True)['pct']
    stats['attendance']=att or 0
    avg=query_db(f'''SELECT ROUND(AVG(m.obtained_marks/m.maximum_marks*100),2) avg FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id JOIN students s ON s.student_id=ss.student_id WHERE {scope}''',sp,fetchone=True)['avg']
    stats['average']=avg or 0
    subject_avgs=query_db(f'''SELECT sub.subject_name label,ROUND(AVG(m.obtained_marks/m.maximum_marks*100),2) value
        FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id
        JOIN students s ON s.student_id=ss.student_id JOIN subjects sub ON sub.subject_id=ss.subject_id
        WHERE {scope} GROUP BY sub.subject_id,sub.subject_name ORDER BY value DESC''',sp)
    recent=query_db(f'''SELECT s.student_name,s.roll_no,sub.subject_name,e.exam_name,sem.semester_name,m.obtained_marks,m.maximum_marks
        FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id JOIN students s ON s.student_id=ss.student_id
        JOIN subjects sub ON sub.subject_id=ss.subject_id JOIN exam_types e ON e.exam_id=m.exam_id JOIN semesters sem ON sem.semester_id=m.semester_id
        WHERE {scope} ORDER BY m.mark_id DESC LIMIT 8''',sp)
    return render_template('index.html',stats=stats,subject_avgs=subject_avgs,recent=recent)

# ---------------- STUDENTS / SUBJECTS ----------------
@app.route('/students')
@login_required
def students():
    q=request.args.get('q','').strip(); scope,sp=student_scope_condition(); cond=f'({scope}) AND (%s="" OR s.roll_no LIKE %s OR s.enrollment_no LIKE %s OR s.student_name LIKE %s OR s.student_class LIKE %s)'; params=sp+[q,f'%{q}%',f'%{q}%',f'%{q}%',f'%{q}%']
    rows=query_db(f'SELECT s.* FROM students s WHERE {cond} ORDER BY s.student_class, CASE WHEN s.roll_no REGEXP "^[0-9]+$" THEN 0 ELSE 1 END, CASE WHEN s.roll_no REGEXP "^[0-9]+$" THEN CAST(s.roll_no AS UNSIGNED) ELSE 0 END,s.roll_no,s.student_name',params)
    return render_template('students.html',students=rows,search=q)

@app.route('/students/add',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def add_student():
    if request.method=='POST':
        roll=request.form.get('roll_no','').strip(); enrollment=request.form.get('enrollment_no','').strip() or roll; name=request.form.get('student_name','').strip(); cls=request.form.get('student_class','').strip()
        if not roll or not name or not cls: flash('Roll number, student name and class are required.','danger')
        elif not enforce_class(cls): flash('You are not assigned to this class.','danger')
        else:
            try: query_db('INSERT INTO students(roll_no,enrollment_no,student_name,student_class) VALUES(%s,%s,%s,%s)',(roll,enrollment,name,cls),commit=True); flash('Student added successfully.','success'); return redirect(url_for('students'))
            except Error: flash('Unable to add student. Roll/enrollment number may already exist.','danger')
    return render_template('add_student.html')

@app.route('/students/<int:student_id>/edit',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def edit_student(student_id):
    s=query_db('SELECT * FROM students WHERE student_id=%s',(student_id,),fetchone=True)
    if not s or not enforce_class(s['student_class']): flash('Student not found or outside your access.','danger'); return redirect(url_for('students'))
    if request.method=='POST':
        try:
            roll=request.form.get('roll_no','').strip(); enrollment=request.form.get('enrollment_no','').strip() or roll; name=request.form.get('student_name','').strip(); cls=request.form.get('student_class','').strip()
            if not enforce_class(cls): raise ValueError('Class is outside your access.')
            remark=request.form.get('teacher_remark','').strip()
            query_db('UPDATE students SET roll_no=%s,enrollment_no=%s,student_name=%s,student_class=%s,teacher_remark=%s WHERE student_id=%s',(roll,enrollment,name,cls,remark,student_id),commit=True); flash('Student updated successfully.','success'); return redirect(url_for('students'))
        except Exception as e: flash(str(e) or 'Unable to update student.','danger')
    return render_template('add_student.html',student=s)

@app.post('/students/<int:student_id>/delete')
@login_required
@role_required('admin','teacher')
def delete_student(student_id):
    s=query_db('SELECT student_class FROM students WHERE student_id=%s',(student_id,),fetchone=True)
    if not s or not enforce_class(s['student_class']):
        flash('You cannot delete this student.','danger')
        return redirect(url_for('students'))
    conn=cursor=None
    try:
        conn=get_db(); cursor=conn.cursor()
        cursor.execute('UPDATE users SET student_id=NULL WHERE student_id=%s',(student_id,))
        cursor.execute('DELETE FROM students WHERE student_id=%s',(student_id,))
        conn.commit()
        flash('Student and all related academic/attendance records were deleted.','success')
    except Exception as exc:
        if conn: conn.rollback()
        app.logger.exception('Student deletion failed: %s',exc)
        flash('Unable to delete the student. No data was changed.','danger')
    finally:
        if cursor: cursor.close()
        if conn: conn.close()
    return redirect(url_for('students'))

@app.post('/students/bulk-delete')
@login_required
@role_required('admin','teacher')
def bulk_delete_students():
    mode=request.form.get('mode','selected')
    ids=[]
    if mode=='selected':
        for raw in request.form.getlist('student_ids'):
            try:
                sid=int(raw)
                if sid not in ids: ids.append(sid)
            except (TypeError,ValueError):
                pass
        if not ids:
            flash('Select at least one student to delete.','warning')
            return redirect(url_for('students',q=request.form.get('q','')))
    else:
        scope,sp=student_scope_condition('s')
        rows=query_db(f'SELECT s.student_id FROM students s WHERE {scope}',sp)
        ids=[int(r['student_id']) for r in rows]
        if not ids:
            flash('There are no students available to delete.','warning')
            return redirect(url_for('students'))

    placeholders=','.join(['%s']*len(ids))
    allowed=query_db(f'SELECT student_id,student_class FROM students WHERE student_id IN ({placeholders})',ids)
    allowed_ids=[int(r['student_id']) for r in allowed if enforce_class(r['student_class'])]
    if len(allowed_ids)!=len(ids):
        flash('One or more selected students are outside your access. Nothing was deleted.','danger')
        return redirect(url_for('students',q=request.form.get('q','')))

    conn=cursor=None
    try:
        conn=get_db(); cursor=conn.cursor()
        ph=','.join(['%s']*len(allowed_ids))
        cursor.execute(f'UPDATE users SET student_id=NULL WHERE student_id IN ({ph})',allowed_ids)
        cursor.execute(f'DELETE FROM students WHERE student_id IN ({ph})',allowed_ids)
        deleted=cursor.rowcount
        conn.commit()
        flash(f'{deleted} student(s) deleted successfully along with related records.','success')
    except Exception as exc:
        if conn: conn.rollback()
        app.logger.exception('Bulk student deletion failed: %s',exc)
        flash('Unable to delete the selected students. No data was changed.','danger')
    finally:
        if cursor: cursor.close()
        if conn: conn.close()
    return redirect(url_for('students',q=request.form.get('q','')))

@app.route('/students/<int:student_id>/subjects',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def student_subjects(student_id):
    student=query_db('SELECT * FROM students WHERE student_id=%s',(student_id,),fetchone=True)
    if not student or not enforce_class(student['student_class']): flash('Student not found or outside your access.','danger'); return redirect(url_for('students'))
    if request.method=='POST':
        ids=set()
        for v in request.form.getlist('subject_ids'):
            try: ids.add(int(v))
            except: pass
        current={x['subject_id'] for x in query_db('SELECT subject_id FROM student_subjects WHERE student_id=%s',(student_id,))}; add=ids-current; rem=current-ids; conn=cur=None
        try:
            conn=get_db(); cur=conn.cursor(dictionary=True)
            for x in add: cur.execute('INSERT INTO student_subjects(student_id,subject_id) VALUES(%s,%s)',(student_id,x))
            if rem:
                ph=','.join(['%s']*len(rem)); cur.execute(f'SELECT COUNT(*) total FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id WHERE ss.student_id=%s AND ss.subject_id IN ({ph})',[student_id,*rem])
                if cur.fetchone()['total']>0: raise ValueError('A subject with existing marks cannot be removed.')
                cur.execute(f'DELETE FROM student_subjects WHERE student_id=%s AND subject_id IN ({ph})',[student_id,*rem])
            conn.commit(); flash('Student subjects saved.','success')
        except Exception as e:
            if conn: conn.rollback()
            flash(str(e),'danger')
        finally:
            if cur: cur.close()
            if conn: conn.close()
        return redirect(url_for('student_subjects',student_id=student_id))
    all_subjects=query_db('SELECT subject_id,subject_name FROM subjects ORDER BY subject_name'); assigned=query_db('SELECT ss.student_subject_id,ss.subject_id,sub.subject_name FROM student_subjects ss JOIN subjects sub ON sub.subject_id=ss.subject_id WHERE ss.student_id=%s ORDER BY sub.subject_name',(student_id,)); assigned_ids={x['subject_id'] for x in assigned}
    return render_template('student_subjects.html',student=student,all_subjects=all_subjects,assigned=assigned,assigned_ids=assigned_ids)

@app.route('/subjects',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def subjects():
    if request.method=='POST':
        name=request.form.get('subject_name','').strip()
        if not name: flash('Subject name is required.','danger')
        else:
            try: query_db('INSERT INTO subjects(subject_name) VALUES(%s)',(name,),commit=True); flash('Subject added.','success')
            except Error: flash('Subject already exists.','danger')
    rows=query_db('''SELECT sub.subject_id,sub.subject_name,COUNT(ss.student_subject_id) assigned_count FROM subjects sub LEFT JOIN student_subjects ss ON ss.subject_id=sub.subject_id GROUP BY sub.subject_id,sub.subject_name ORDER BY sub.subject_name''')
    return render_template('subjects.html',subjects=rows)

@app.route('/subjects/<int:subject_id>/edit',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def edit_subject(subject_id):
    subject=query_db('SELECT * FROM subjects WHERE subject_id=%s',(subject_id,),fetchone=True)
    if not subject: flash('Subject not found.','danger'); return redirect(url_for('subjects'))
    if request.method=='POST':
        try: query_db('UPDATE subjects SET subject_name=%s WHERE subject_id=%s',(request.form.get('subject_name','').strip(),subject_id),commit=True); flash('Subject updated.','success'); return redirect(url_for('subjects'))
        except Error: flash('Unable to update subject.','danger')
    return render_template('edit_subject.html',subject=subject)

@app.post('/subjects/<int:subject_id>/delete')
@login_required
@role_required('admin','teacher')
def delete_subject(subject_id):
    conn=None; cur=None
    try:
        conn=get_db(); cur=conn.cursor(dictionary=True); conn.start_transaction()
        cur.execute('SELECT student_subject_id FROM student_subjects WHERE subject_id=%s',(subject_id,))
        ids=[r['student_subject_id'] for r in cur.fetchall()]
        if ids:
            marks_ph=','.join(['%s']*len(ids))
            cur.execute(f'DELETE FROM marks WHERE student_subject_id IN ({marks_ph})',ids)
            cur.execute(f'DELETE FROM student_subjects WHERE student_subject_id IN ({marks_ph})',ids)
        cur.execute('DELETE FROM subjects WHERE subject_id=%s',(subject_id,))
        conn.commit(); flash('Subject and its related marks/assignments were deleted.','success')
    except Exception as e:
        if conn: conn.rollback()
        flash('Unable to delete subject: '+str(e),'danger')
    finally:
        if cur: cur.close()
        if conn: conn.close()
    return redirect(url_for('subjects'))

@app.post('/subjects/delete-all')
@login_required
@role_required('admin','teacher')
def delete_all_subjects():
    conn=None; cur=None
    try:
        conn=get_db(); cur=conn.cursor(); conn.start_transaction()
        cur.execute('DELETE FROM marks')
        cur.execute('DELETE FROM student_subjects')
        cur.execute('DELETE FROM subjects')
        conn.commit(); flash('All subjects, assignments and related marks were deleted. Students and attendance were preserved.','success')
    except Exception as e:
        if conn: conn.rollback()
        flash('Unable to delete all subjects: '+str(e),'danger')
    finally:
        if cur: cur.close()
        if conn: conn.close()
    return redirect(url_for('subjects'))

@app.post('/student-subject/<int:student_subject_id>/delete')
@login_required
@role_required('admin','teacher')
def delete_student_subject(student_subject_id):
    try: query_db('DELETE FROM student_subjects WHERE student_subject_id=%s',(student_subject_id,),commit=True); flash('Subject unassigned.','success')
    except Error: flash('Subject has marks and cannot be removed.','danger')
    return redirect(url_for('students'))

# ---------------- SETTINGS / ASSESSMENTS ----------------
@app.route('/settings',methods=['GET','POST'])
@login_required
@role_required('admin')
def settings():
    if request.method=='POST':
        typ=request.form.get('institution_type','college'); name=request.form.get('institution_name','').strip() or 'Academic Marks Management System'; use='1' if request.form.get('use_enrollment')=='1' else '0'; threshold=safe_float(request.form.get('detention_threshold'))
        threshold=threshold if threshold is not None and 0<=threshold<=100 else 75
        for k,v in [('institution_type',typ if typ in ('school','college') else 'college'),('institution_name',name),('use_enrollment_no',use),('detention_threshold',str(threshold))]: query_db('UPDATE settings SET setting_value=%s WHERE setting_key=%s',(v,k),commit=True)
        flash('Settings saved.','success'); return redirect(url_for('settings'))
    return render_template('settings.html',config=institution_config(),semesters=query_db('SELECT * FROM semesters ORDER BY semester_id'),exams=query_db('SELECT * FROM exam_types ORDER BY exam_id'))

@app.post('/settings/assessment')
@login_required
@role_required('admin','teacher')
def assessment_setting():
    kind=request.form.get('kind'); name=request.form.get('name','').strip(); table='semesters' if kind=='semester' else 'exam_types'; col='semester_name' if kind=='semester' else 'exam_name'
    if not name: flash('Name is required.','danger')
    else:
        try: query_db(f'INSERT INTO {table}({col}) VALUES(%s)',(name,),commit=True); flash('Assessment type added.','success')
        except Error: flash('This assessment type already exists.','danger')
    return redirect(url_for('settings'))

@app.post('/settings/delete/<kind>/<int:item_id>')
@login_required
@role_required('admin')
def delete_setting(kind,item_id):
    conn=None; cur=None
    try:
        conn=get_db(); cur=conn.cursor(); conn.start_transaction()
        if kind=='exam':
            cur.execute('DELETE FROM marks WHERE exam_id=%s',(item_id,))
            cur.execute('DELETE FROM exam_types WHERE exam_id=%s',(item_id,))
        elif kind=='semester':
            cur.execute('DELETE FROM marks WHERE semester_id=%s',(item_id,))
            cur.execute('DELETE FROM semesters WHERE semester_id=%s',(item_id,))
        else:
            raise ValueError('Invalid setting type')
        conn.commit(); flash('Academic setting and its related marks were deleted.','success')
    except Exception as e:
        if conn: conn.rollback()
        flash('Unable to delete setting: '+str(e),'danger')
    finally:
        if cur: cur.close()
        if conn: conn.close()
    return redirect(url_for('settings'))

@app.post('/settings/delete-all-exams')
@login_required
@role_required('admin')
def delete_all_exams():
    conn=None; cur=None
    try:
        conn=get_db(); cur=conn.cursor(); conn.start_transaction()
        cur.execute('DELETE FROM marks')
        cur.execute('DELETE FROM exam_types')
        conn.commit(); flash('All test/assessment types and related marks were deleted. Students and subjects were preserved.','success')
    except Exception as e:
        if conn: conn.rollback()
        flash('Unable to delete all test types: '+str(e),'danger')
    finally:
        if cur: cur.close()
        if conn: conn.close()
    return redirect(url_for('settings'))

# ---------------- MARKS ENTRY ----------------
@app.route('/marks',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def marks():
    scope,sp=student_scope_condition(); students=query_db(f'SELECT s.* FROM students s WHERE {scope} ORDER BY s.student_class,s.student_name',sp); semesters=query_db('SELECT * FROM semesters ORDER BY semester_id'); exams=query_db('SELECT * FROM exam_types ORDER BY exam_id')
    if request.method=='POST':
        sid=request.form.get('student_id',type=int); sem=request.form.get('semester_id',type=int); exam=request.form.get('exam_id',type=int)
        st=query_db('SELECT * FROM students WHERE student_id=%s',(sid,),fetchone=True)
        if not st or not enforce_class(st['student_class']): flash('Student is outside your access.','danger'); return redirect(url_for('marks'))
        assignments=query_db('SELECT ss.student_subject_id,sub.subject_name FROM student_subjects ss JOIN subjects sub ON sub.subject_id=ss.subject_id WHERE ss.student_id=%s',(sid,)); saved=0; errors=[]
        for a in assignments:
            o=request.form.get(f'obtained_{a["student_subject_id"]}','').strip(); mx=request.form.get(f'maximum_{a["student_subject_id"]}','').strip()
            if o=='' and mx=='': continue
            of=safe_float(o); mf=safe_float(mx)
            if of is None or mf is None or mf<=0 or of<0 or of>mf: errors.append(a['subject_name']+': invalid marks'); continue
            try: query_db('INSERT INTO marks(student_subject_id,semester_id,exam_id,obtained_marks,maximum_marks) VALUES(%s,%s,%s,%s,%s) ON DUPLICATE KEY UPDATE obtained_marks=VALUES(obtained_marks),maximum_marks=VALUES(maximum_marks)',(a['student_subject_id'],sem,exam,of,mf),commit=True); saved+=1
            except Error: errors.append(a['subject_name']+': save failed')
        if saved: flash(f'{saved} subject mark(s) saved.','success')
        if errors: flash(' | '.join(errors),'danger')
        return redirect(url_for('marks',student_id=sid,semester_id=sem,exam_id=exam))
    sid=request.args.get('student_id',type=int); sem=request.args.get('semester_id',type=int); exam=request.args.get('exam_id',type=int); student=query_db('SELECT * FROM students WHERE student_id=%s',(sid,),fetchone=True) if sid else None
    assignments=query_db('SELECT ss.student_subject_id,sub.subject_name FROM student_subjects ss JOIN subjects sub ON sub.subject_id=ss.subject_id WHERE ss.student_id=%s ORDER BY sub.subject_name',(sid,)) if sid else []
    existing={}
    if sid and sem and exam:
        rows=query_db('SELECT student_subject_id,obtained_marks,maximum_marks FROM marks WHERE student_subject_id IN (SELECT student_subject_id FROM student_subjects WHERE student_id=%s) AND semester_id=%s AND exam_id=%s',(sid,sem,exam)); existing={r['student_subject_id']:r for r in rows}
    return render_template('marks.html',students=students,semesters=semesters,exams=exams,selected_student=sid,selected_semester=sem,selected_exam=exam,student=student,assignments=assignments,existing=existing)

@app.route('/marks/<int:mark_id>/edit',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def edit_marks(mark_id):
    r=query_db('''SELECT m.*,s.student_name,s.student_class,sub.subject_name,sem.semester_name,e.exam_name FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id JOIN students s ON s.student_id=ss.student_id JOIN subjects sub ON sub.subject_id=ss.subject_id JOIN semesters sem ON sem.semester_id=m.semester_id JOIN exam_types e ON e.exam_id=m.exam_id WHERE m.mark_id=%s''',(mark_id,),fetchone=True)
    if not r or not enforce_class(r['student_class']): flash('Marks record not found or outside your access.','danger'); return redirect(url_for('marks'))
    if request.method=='POST':
        o=safe_float(request.form.get('obtained_marks')); mx=safe_float(request.form.get('maximum_marks'))
        if o is None or mx is None or mx<=0 or o<0 or o>mx: flash('Enter valid marks. Obtained cannot exceed maximum.','danger')
        else:
            query_db('UPDATE marks SET obtained_marks=%s,maximum_marks=%s WHERE mark_id=%s',(o,mx,mark_id),commit=True); flash('Marks updated successfully.','success'); return redirect(url_for('marks'))
    return render_template('edit_marks.html',record=r)

@app.post('/marks/<int:mark_id>/delete')
@login_required
@role_required('admin','teacher')
def delete_marks(mark_id):
    r=query_db('SELECT s.student_class FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id JOIN students s ON s.student_id=ss.student_id WHERE m.mark_id=%s',(mark_id,),fetchone=True)
    if r and enforce_class(r['student_class']): query_db('DELETE FROM marks WHERE mark_id=%s',(mark_id,),commit=True); flash('Marks deleted successfully.','success')
    else: flash('You cannot delete this mark.','danger')
    return redirect(url_for('marks'))

# ---------------- FILE IMPORT ----------------
def cleanup_import_files(max_age_hours=12):
    now=datetime.now().timestamp(); cutoff=now-(max_age_hours*3600)
    try:
        for path in Path(IMPORT_DIR).glob('*.json'):
            try:
                if path.stat().st_mtime < cutoff: path.unlink()
            except OSError: pass
        for d in Path(IMPORT_DIR).glob('*'):
            if d.is_dir():
                try:
                    if d.stat().st_mtime < cutoff:
                        import shutil as _shutil; _shutil.rmtree(d,ignore_errors=True)
                except OSError: pass
    except OSError: pass

cleanup_import_files()
HEADER_MAP={'rollno':'roll_no','rollnumber':'roll_no','roll':'roll_no','enrollmentno':'enrollment_no','enrollmentnumber':'enrollment_no','enrollment':'enrollment_no','studentname':'student_name','name':'student_name','student':'student_name','studentclass':'student_class','class':'student_class','classdivision':'student_class','division':'student_class','subject':'subject_name','subjectname':'subject_name','course':'subject_name','semester':'semester_name','term':'semester_name','exam':'exam_name','examtype':'exam_name','examname':'exam_name','assessment':'exam_name','assessmentname':'exam_name','test':'exam_name','testname':'exam_name','testtype':'exam_name','obtained':'obtained_marks','obtainedmarks':'obtained_marks','marks':'obtained_marks','score':'obtained_marks','mark':'obtained_marks','maximum':'maximum_marks','maxmarks':'maximum_marks','maximummarks':'maximum_marks','totalmarks':'maximum_marks','outof':'maximum_marks'}
ATT_MAP={'rollno':'roll_no','rollnumber':'roll_no','roll':'roll_no','enrollmentno':'enrollment_no','enrollmentnumber':'enrollment_no','enrollment':'enrollment_no','studentname':'student_name','name':'student_name','student':'student_name','studentclass':'student_class','class':'student_class','subject':'subject_name','subjectname':'subject_name','month':'month','monthname':'month','monthyear':'month','workingdays':'working_days','workingday':'working_days','totaldays':'working_days','presentdays':'present_days','present':'present_days','dayspresent':'present_days','attendance':'attendance_pct','percentage':'attendance_pct','attendancepercentage':'attendance_pct','notes':'notes'}
def clean_header(x): return re.sub(r'[^a-z0-9]','',str(x or '').lower().strip())

def _is_blank_header(header):
    """Return True for columns that are genuinely blank/placeholder headers.
    Spreadsheet libraries often expose empty headers as None, whitespace,
    'Unnamed: n', or generated 'Column n' labels. These columns are ignored
    by the importer rather than causing a hard failure.
    """
    if header is None:
        return True
    h=str(header).strip()
    if not h:
        return True
    return bool(re.fullmatch(r'(?i)(unnamed\s*:?\s*\d+|column\s+\d+)', h))

def extract_mark_and_max(value):
    """Extract obtained/max marks from flexible cell values.
    Accepts values such as 85, '85/100', '85 / 100', '85 out of 100',
    '85 (100)', and Excel numeric values. Returns strings for preview/editing.
    """
    if value is None:
        return '', ''
    if isinstance(value, bool):
        return '', ''
    text=str(value).strip()
    if not text:
        return '', ''
    # Common obtained/max formats.
    m=re.search(r'(-?\d+(?:\.\d+)?)\s*(?:/|\\|\bof\b|out\s+of)\s*(-?\d+(?:\.\d+)?)', text, re.I)
    if m:
        return m.group(1), m.group(2)
    m=re.search(r'(-?\d+(?:\.\d+)?)\s*\(\s*(-?\d+(?:\.\d+)?)\s*\)', text)
    if m:
        return m.group(1), m.group(2)
    m=re.search(r'-?\d+(?:\.\d+)?', text)
    return (m.group(0), '') if m else ('', '')

def _subject_candidate(header, raw_rows):
    """Heuristic for a horizontal subject/marks column.
    Metadata/assessment columns are excluded. A column qualifies when it has
    at least one numeric mark-like cell; this keeps subject names flexible.
    """
    h=str(header or '').strip()
    c=clean_header(h)
    if not h or _is_blank_header(h) or c in HEADER_MAP:
        return False
    if re.search(r'(?i)\b(unit\s*\d+|ut\s*\d+|class\s*test|periodic|quiz|mid\s*term|midterm|semester|sem\s*\d+|final|annual|internal|assessment|test\s*\d+)\b',h):
        return False
    # Exclude obvious maximum/percentage/helper columns.
    if re.search(r'(?i)\b(max|maximum|total|percentage|percent|grade|remark|note|attendance)\b',h):
        return False
    numeric=0
    nonempty=0
    for row in raw_rows[:1000]:
        v=row.get(header,'')
        if v is None or not str(v).strip():
            continue
        nonempty += 1
        mark,_=extract_mark_and_max(v)
        if mark != '':
            try:
                float(mark); numeric += 1
            except ValueError:
                pass
    return nonempty > 0 and numeric > 0
def rows_from_matrix(values):
    """Read a table even when the title/instructions appear above the header.
    The best header row is detected from known academic column names, so row 1
    does not have to be the header and column order is irrelevant.
    """
    matrix=[list(r) for r in values if any(str(x or '').strip() for x in r)]
    if not matrix: return []
    known=set(HEADER_MAP) | {'rollno','rollnumber','roll','name','student','class','subject','marks','score','outof','maximum','maxmarks','test','assessment','exam','semester','term'}
    best_idx=0; best_score=-1
    for idx,row in enumerate(matrix[:30]):
        score=sum(1 for x in row if clean_header(x) in known)
        if score>best_score:
            best_idx,best_score=idx,score
    # If no recognizable header exists, preserve the old behavior.
    if best_score<=0: best_idx=0
    header=matrix[best_idx]
    heads=[]; seen={}
    for i,x in enumerate(header):
        h=str(x or '').strip() or f'Column {i+1}'; base=h; n=2
        while h in seen: h=f'{base} {n}'; n+=1
        seen[h]=1; heads.append(h)
    out=[]
    for row in matrix[best_idx+1:]:
        vals=list(row)+['']*(len(heads)-len(row))
        d=dict(zip(heads,vals[:len(heads)]))
        # Ignore repeated header rows and completely empty rows.
        if clean_header(' '.join(str(x or '') for x in row)) and any(clean_header(str(x or '')) in known for x in row):
            if sum(1 for x in row if clean_header(x) in known)>=2: continue
        if any(str(x or '').strip() for x in row): out.append(d)
    return out

def read_text(raw):
    text=raw.decode('utf-8-sig',errors='replace'); sample='\n'.join(text.splitlines()[:30])
    try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t|')
    except csv.Error: dialect=csv.excel_tab if '\t' in sample else csv.excel
    # Read as a matrix so the header detector can ignore title rows.
    return rows_from_matrix(csv.reader(io.StringIO(text),dialect=dialect))

def read_file(f):
    """Read academic tabular files using disk-backed parsing for large uploads.
    Supported: XLSX, XLSM, XLTX, XLTM, XLS, CSV, TSV, TXT and JSON.
    Excel worksheets are read in openpyxl read-only mode where possible.
    """
    filename = getattr(f, 'filename', '') or ''
    path = f if isinstance(f, (str, os.PathLike)) else None
    ext = Path(filename if path is None else str(path)).suffix.lower()

    if path is None:
        raw = f.read()
        if not raw:
            return []
        if ext == '.json' or raw.lstrip().startswith((b'{', b'[')):
            data = json.loads(raw.decode('utf-8-sig', errors='replace'))
            data = data.get('data', data.get('rows', [data])) if isinstance(data, dict) else data
            return data if isinstance(data, list) and (not data or isinstance(data[0], dict)) else rows_from_matrix(data)
        if ext in ('.xlsx', '.xlsm', '.xltx', '.xltm') or raw[:2] == b'PK':
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            result = []
            try:
                for ws in wb.worksheets:
                    result.extend(rows_from_matrix(ws.iter_rows(values_only=True)))
                return result
            finally:
                wb.close()
        if ext == '.xls' or raw[:8] == bytes.fromhex('D0CF11E0A1B11AE1'):
            if not xlrd:
                raise ValueError('Install xlrd for legacy .xls files.')
            book = xlrd.open_workbook(file_contents=raw, on_demand=True)
            result = []
            try:
                for sh in book.sheets():
                    result.extend(rows_from_matrix(sh.row_values(i) for i in range(sh.nrows)))
                return result
            finally:
                book.release_resources()
        if b'\x00' not in raw[:4096]:
            return read_text(raw)
        raise ValueError('The file is not a readable tabular file. Supported formats: XLSX, XLSM, XLTX, XLTM, XLS, CSV, TSV, TXT and JSON.')

    path = str(path)
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        return []

    if ext == '.json':
        with open(path, 'rb') as fh:
            raw = fh.read()
        data = json.loads(raw.decode('utf-8-sig', errors='replace'))
        data = data.get('data', data.get('rows', [data])) if isinstance(data, dict) else data
        return data if isinstance(data, list) and (not data or isinstance(data[0], dict)) else rows_from_matrix(data)

    if ext in ('.xlsx', '.xlsm', '.xltx', '.xltm'):
        wb = load_workbook(path, data_only=True, read_only=True)
        result = []
        try:
            for ws in wb.worksheets:
                result.extend(rows_from_matrix(ws.iter_rows(values_only=True)))
            return result
        finally:
            wb.close()

    if ext == '.xls':
        if not xlrd:
            raise ValueError('Install xlrd for legacy .xls files.')
        book = xlrd.open_workbook(path, on_demand=True)
        result = []
        try:
            for sh in book.sheets():
                result.extend(rows_from_matrix(sh.row_values(i) for i in range(sh.nrows)))
            return result
        finally:
            book.release_resources()

    if ext in ('.csv', '.tsv', '.txt') or ext == '':
        with open(path, 'rb') as fh:
            raw = fh.read()
        if b'\x00' not in raw[:4096]:
            return read_text(raw)

    raise ValueError('The file is not a readable tabular file. Supported formats: XLSX, XLSM, XLTX, XLTM, XLS, CSV, TSV, TXT and JSON.')

def normalize(rows,map_):
    if not rows: return []
    heads=list(rows[0].keys()); canonical={h:map_.get(clean_header(h)) for h in heads}; fields=list(dict.fromkeys(map_.values())); out=[]
    for n,raw in enumerate(rows,2):
        r={'source_row':n}
        for field in fields:
            r[field]=''
            for h,c in canonical.items():
                if c==field: r[field]=str(raw.get(h,'') if raw.get(h,'') is not None else '').strip(); break
        if any(v for k,v in r.items() if k!='source_row'): out.append(r)
    return out

def assessment_name_from_header(header):
    h=str(header or '').strip()
    h=re.sub(r'(?i)(?:marks?|score|obtained|maximum|max|out\s*of|total)\s*$', '', h).strip(' -_:()')
    return h or str(header or '').strip()

def parse_assessment_subject_header(header):
    """Split headers such as 'Unit 1 Mathematics', 'Mathematics - Unit 2',
    'Semester 1 | Physics' into (assessment, subject). If a subject cannot
    be confidently separated, subject is returned blank and the row-level
    Subject column can still supply it.
    """
    h=re.sub(r'(?i)(?:marks?|score|obtained|maximum|max|out\s*of|total)\s*$', '', str(header or '').strip()).strip(' -_:|()')
    patterns=[
        r'^(unit\s*\d+|ut\s*\d+|class\s*test\s*\d+|periodic\s*test\s*\d+|quiz\s*\d+|mid\s*term|midterm|semester\s*\d*|sem\s*\d+|final|annual|internal|assessment\s*\d+|test\s*\d+)[\s\-_:|/]+(.+)$',
        r'^(.+?)[\s\-_:|/]+(unit\s*\d+|ut\s*\d+|class\s*test\s*\d+|periodic\s*test\s*\d+|quiz\s*\d+|mid\s*term|midterm|semester\s*\d*|sem\s*\d+|final|annual|internal|assessment\s*\d+|test\s*\d+)$'
    ]
    for pat in patterns:
        m=re.match(pat,h,re.I)
        if m:
            if re.match(pat, h, re.I).group(1).lower().strip() in ('semester','sem'):
                pass
            a,b=m.group(1).strip(' -_:|'),m.group(2).strip(' -_:|')
            # In reverse pattern the assessment is group 2.
            if re.search(r'(?i)^(unit|ut|class\s*test|periodic|quiz|mid|semester|sem|final|annual|internal|assessment|test)', b):
                return b,a
            return a,b
    return assessment_name_from_header(h), ''

def _infer_semester_from_exam(exam_name):
    m=re.search(r'(?i)\b(?:semester|sem)\s*(\d+)\b',str(exam_name or ''))
    return f'Semester {m.group(1)}' if m else ''

def normalize_marks_rows(raw_rows, selected_exam_name=''):
    """College-scale, order-independent marks importer.
    Supports long and wide layouts, multiple exams in one file, horizontal
    subjects, multiple Excel sheets, blank columns, title rows, and unknown
    assessment names (unknown types are created during save).
    """
    if not raw_rows: return []
    raw_rows=[{k:v for k,v in raw.items() if not _is_blank_header(k)} for raw in raw_rows]
    rows=normalize(raw_rows,HEADER_MAP)
    if not rows: return []
    raw_heads=list(raw_rows[0].keys())

    # 1) If an explicit Exam/Test column exists, preserve every assessment.
    explicit_exam=any(str(r.get('exam_name','')).strip() for r in rows)
    out=[]
    if explicit_exam:
        for raw,r in zip(raw_rows,rows):
            r['exam_name']=r.get('exam_name','').strip() or selected_exam_name
            # Normal long layout.
            if r.get('subject_name') and str(r.get('obtained_marks','')).strip():
                out.append(r); continue
            # Exam column + subjects horizontally.
            subject_cols=[h for h in raw_heads if _subject_candidate(h,raw_rows)]
            for sh in subject_cols:
                mark,mx=extract_mark_and_max(raw.get(sh,''))
                if not str(mark).strip(): continue
                nr=dict(r); nr['subject_name']=str(sh).strip(); nr['obtained_marks']=mark; nr['maximum_marks']=mx or '100'
                out.append(nr)
        return out

    # 2) Detect wide assessment columns. Do not treat metadata as assessments.
    assessment_cols=[]
    for h in raw_heads:
        c=clean_header(h)
        if c in HEADER_MAP: continue
        if re.search(r'(?i)\b(unit\s*\d+|ut\s*\d+|class\s*test|periodic|quiz|mid\s*term|midterm|semester|sem\s*\d+|final|annual|internal|assessment|test\s*\d+)\b',str(h)):
            assessment_cols.append(h)

    # 3) Horizontal subject columns with one default/selected assessment.
    subject_cols=[h for h in raw_heads if _subject_candidate(h,raw_rows)]
    if not assessment_cols and subject_cols:
        for raw,r in zip(raw_rows,rows):
            base=dict(r); base['exam_name']=r.get('exam_name','').strip() or selected_exam_name
            for sh in subject_cols:
                mark,mx=extract_mark_and_max(raw.get(sh,''))
                if not str(mark).strip(): continue
                nr=dict(base); nr['subject_name']=str(sh).strip(); nr['obtained_marks']=mark; nr['maximum_marks']=mx or '100'
                out.append(nr)
        return out

    # 4) Wide assessment columns: each assessment column becomes a separate
    # mark record. If the header contains the subject, derive it too.
    max_headers={clean_header(h):h for h in raw_heads if re.search(r'(?i)(max|maximum|out\s*of)',str(h))}
    for raw,r in zip(raw_rows,rows):
        for ah in assessment_cols:
            val=raw.get(ah,'')
            mark,mx=extract_mark_and_max(val)
            if not str(mark).strip(): continue
            exam_name,header_subject=parse_assessment_subject_header(ah)
            # Prefer a row-level Subject column when present.
            subject_name=r.get('subject_name','').strip() or header_subject
            if not mx:
                target=clean_header(ah)
                for mh,morig in max_headers.items():
                    if target in mh or mh in target or target.replace('marks','') in mh:
                        mx=str(raw.get(morig,'') or '').strip(); break
            nr={k:r.get(k,'') for k in ('roll_no','enrollment_no','student_name','student_class','subject_name','semester_name')}
            nr['subject_name']=subject_name
            nr['exam_name']=exam_name or selected_exam_name
            nr['semester_name']=nr['semester_name'] or _infer_semester_from_exam(nr['exam_name'])
            nr['obtained_marks']=mark; nr['maximum_marks']=mx or '100'
            # If assessment header did not contain a subject, a row-level
            # subject is mandatory; otherwise it will be flagged in preview.
            out.append(nr)

    # 5) Fallback for a sheet that has a Test/Assessment title row plus
    # horizontal subjects but did not match a standard assessment header.
    if not out and subject_cols:
        for raw,r in zip(raw_rows,rows):
            for sh in subject_cols:
                mark,mx=extract_mark_and_max(raw.get(sh,''))
                if not str(mark).strip(): continue
                nr={k:r.get(k,'') for k in ('roll_no','enrollment_no','student_name','student_class','semester_name','exam_name')}
                nr['exam_name']=nr.get('exam_name') or selected_exam_name
                nr['subject_name']=str(sh).strip(); nr['obtained_marks']=mark; nr['maximum_marks']=mx or '100'
                nr['semester_name']=nr['semester_name'] or _infer_semester_from_exam(nr['exam_name'])
                out.append(nr)
    return out

def exam_choices(): return query_db('SELECT exam_id,exam_name FROM exam_types ORDER BY exam_id')
def semester_choices(): return query_db('SELECT semester_id,semester_name FROM semesters ORDER BY semester_id')

@app.route('/marks/import',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def import_marks():
    if request.method=='POST':
        f=request.files.get('marks_file'); sem=request.form.get('semester_id',type=int)
        exam=request.form.get('exam_id',type=int); custom=request.form.get('custom_exam','').strip(); import_mode=request.form.get('import_mode','overall')
        if not f or not f.filename:
            flash('Select a marks file first.','danger'); return redirect(url_for('import_marks'))
        try:
            # A selected/custom assessment is only a fallback. If the file has
            # its own Test/Exam/Assessment column or wide Unit/Sem columns,
            # every assessment in the file is retained.
            default_exam_name=''
            # Overall/auto-detect deliberately leaves the default assessment empty.
            # The importer will preserve every Unit/Sem/Test found in the file.
            if import_mode != 'overall' and custom:
                ex=query_db('SELECT exam_id,exam_name FROM exam_types WHERE LOWER(exam_name)=LOWER(%s)',(custom,),fetchone=True)
                if ex: exam=ex['exam_id']; default_exam_name=ex['exam_name']
                else:
                    exam=query_db('INSERT INTO exam_types(exam_name) VALUES(%s)',(custom,),commit=True); default_exam_name=custom
            elif import_mode != 'overall' and exam:
                ex=query_db('SELECT exam_name FROM exam_types WHERE exam_id=%s',(exam,),fetchone=True)
                default_exam_name=ex['exam_name'] if ex else ''

            # Save the upload to disk with its real extension before parsing.
            # This keeps the previous upload -> preview flow while allowing large
            # files to be parsed from disk instead of loading the upload object
            # through a background job.
            upload_ext=Path(f.filename).suffix.lower() or '.upload'
            temp_path=os.path.join(IMPORT_DIR, uuid.uuid4().hex + upload_ext)
            try:
                f.save(temp_path)
                raw_rows=read_file(temp_path)
            finally:
                try:
                    os.remove(temp_path)
                except OSError:
                    pass
            rows=normalize_marks_rows(raw_rows,default_exam_name)
            validate_import_row_count(len(rows))
            if not rows:
                flash('No readable marks data was found. The system can handle columns in any order and title rows above the table, but the file must contain student, subject and marks information.','danger')
                return redirect(url_for('import_marks'))

            # Determine whether the file itself contains assessment information.
            has_assessments=any(str(r.get('exam_name','')).strip() for r in rows)
            if not has_assessments:
                flash('The file does not contain a test/assessment name. Select a default test above or add a Test/Exam/Assessment column to the file.','danger')
                return redirect(url_for('import_marks'))
            if not sem and not any(str(r.get('semester_name','')).strip() for r in rows):
                # Unit/assessment-only files can still be imported; the teacher can
                # assign/edit the semester in the review table before saving.
                for r in rows:
                    r['semester_name']=_infer_semester_from_exam(r.get('exam_name','')) or 'Semester 1'

            token=uuid.uuid4().hex
            payload={'kind':'marks','semester_id':sem,'exam_id':exam,'default_exam_name':default_exam_name,'rows':rows,'source_filename':f.filename,'import_mode':import_mode}
            with open(os.path.join(IMPORT_DIR,token+'.json'),'w',encoding='utf-8') as out:
                json.dump(payload,out,ensure_ascii=False)
            session['import_token']=token
            return redirect(url_for('import_preview'))
        except Exception as exc:
            app.logger.exception('Marks import read failed: %s',exc)
            msg=str(exc)
            if 'column' in msg.lower() and 'empty' in msg.lower():
                msg='Blank/empty columns were found. They are ignored automatically; no special column sequence is required.'
            flash(f'File could not be read: {msg}','danger')
    return render_template('import_marks.html',semesters=semester_choices(),exams=exam_choices())

def _validate_marks_rows(edited, payload):
    fields=['roll_no','enrollment_no','student_name','student_class','subject_name','semester_name','exam_name','obtained_marks','maximum_marks']
    clean=[]; errors=[]; seen={}; roll_en={}; en_roll={}
    semester_default=''
    if payload.get('semester_id'):
        x=query_db('SELECT semester_name FROM semesters WHERE semester_id=%s',(payload['semester_id'],),fetchone=True); semester_default=x['semester_name'] if x else ''
    exam_default=payload.get('default_exam_name','')
    for i,item in enumerate(edited,1):
        r={k:str(item.get(k,'') or '').strip() for k in fields}
        if not any(r.values()): continue
        r['enrollment_no']=r['enrollment_no'] or r['roll_no']
        r['semester_name']=r['semester_name'] or semester_default
        r['exam_name']=r['exam_name'] or exam_default
        if not r['roll_no'] or not r['student_name'] or not r['subject_name']:
            errors.append(f'Row {i}: Roll No, Student Name and Subject are required.'); continue
        if not r['semester_name'] or not r['exam_name']:
            errors.append(f'Row {i}: Semester/Term and Test/Assessment are required.'); continue
        o=safe_float(r['obtained_marks']); mx=safe_float(r['maximum_marks'])
        if o is None or mx is None or mx<=0 or o<0 or o>mx:
            errors.append(f'Row {i}: invalid marks.'); continue
        if not enforce_class(r['student_class'] or 'Not Provided') and current_user()['role']!='admin':
            errors.append(f'Row {i}: class is outside your assigned classes.'); continue
        rk=r['roll_no'].casefold(); ek=r['enrollment_no'].casefold()
        if rk in roll_en and roll_en[rk]!=ek: errors.append(f'Row {i}: one Roll No has multiple Enrollment Nos.'); continue
        if ek in en_roll and en_roll[ek]!=rk: errors.append(f'Row {i}: one Enrollment No has multiple Roll Nos.'); continue
        # A marks import is an upsert. If the same student + subject +
        # semester + assessment appears more than once in the uploaded file,
        # keep the latest row instead of rejecting the entire import. This is
        # important for real college exports where repeated headers/sections
        # or edited rows can contain the same logical record more than once.
        key=(rk,r['subject_name'].casefold(),r['semester_name'].casefold(),r['exam_name'].casefold())
        if key in seen:
            clean[seen[key]]=r
            roll_en[rk]=ek; en_roll[ek]=rk
            continue
        roll_en[rk]=ek; en_roll[ek]=rk; seen[key]=len(clean); clean.append(r)
    return clean, errors

def _chunks(seq, size=IMPORT_BATCH_SIZE):
    seq=list(seq)
    for i in range(0,len(seq),size): yield seq[i:i+size]

def _select_by_in(cur, sql_prefix, values, sql_suffix=''):
    out=[]
    for batch in _chunks(list(dict.fromkeys(values))):
        ph=','.join(['%s']*len(batch)); cur.execute(sql_prefix.format(ph=ph)+sql_suffix,batch); out.extend(cur.fetchall())
    return out

def _executemany_chunks(cur, sql, rows):
    rows=list(rows)
    for batch in _chunks(rows):
        cur.executemany(sql,batch)

def _fetch_map(cur, table, id_col, name_col, names):
    names=[str(x).strip() for x in names if str(x).strip()]
    if not names: return {}
    unique=list(dict.fromkeys(names))
    for batch in _chunks(unique):
        cur.executemany(f'INSERT IGNORE INTO {table}({name_col}) VALUES(%s)', [(n,) for n in batch])
    rows=[]
    for batch in _chunks(unique):
        ph=','.join(['%s']*len(batch)); cur.execute(f'SELECT {id_col},{name_col} FROM {table} WHERE LOWER({name_col}) IN ({ph})',[n.lower() for n in batch]); rows.extend(cur.fetchall())
    return {str(r[name_col]).casefold():r[id_col] for r in rows}

def _save_marks_rows(payload, edited):
    clean, errors=_validate_marks_rows(edited,payload)
    if errors: raise ValueError(' | '.join(errors[:12]))
    if not clean: return 0
    conn=cur=None
    try:
        conn=get_db(); conn.start_transaction(); cur=conn.cursor(dictionary=True)
        role=current_user()['role']
        # Validate/resolve students in batches instead of 2 SELECTs per row.
        rolls=list(dict.fromkeys(r['roll_no'] for r in clean)); ens=list(dict.fromkeys(r['enrollment_no'] or r['roll_no'] for r in clean))
        by_roll={str(x['roll_no']).casefold():x for x in _select_by_in(cur,'SELECT * FROM students WHERE roll_no IN ({ph})',rolls)}
        by_en={str(x['enrollment_no']).casefold():x for x in _select_by_in(cur,'SELECT * FROM students WHERE enrollment_no IN ({ph})',ens) if x.get('enrollment_no')}
        student_by_key={}
        for r in clean:
            roll=r['roll_no']; en=r['enrollment_no'] or roll; a=by_roll.get(roll.casefold()); b=by_en.get(en.casefold())
            if a and b and a['student_id']!=b['student_id']:
                raise ValueError(f'Roll No {roll} and Enrollment No {en} belong to different students.')
            stu=a or b
            cls=r['student_class'] or (stu['student_class'] if stu else 'Not Provided')
            if not enforce_class(cls) and role!='admin': raise ValueError(f'Class {cls} is outside your assigned classes.')
            student_by_key[(roll.casefold(),en.casefold())]=(stu,cls)
        # Upsert students in one batch. The database unique keys protect duplicates.
        student_rows=[]
        for r in clean:
            stu,cls=student_by_key[(r['roll_no'].casefold(),(r['enrollment_no'] or r['roll_no']).casefold())]
            student_rows.append((r['roll_no'],r['enrollment_no'] or r['roll_no'],r['student_name'],cls))
        _executemany_chunks(cur,'INSERT INTO students(roll_no,enrollment_no,student_name,student_class) VALUES(%s,%s,%s,%s) ON DUPLICATE KEY UPDATE enrollment_no=VALUES(enrollment_no),student_name=VALUES(student_name),student_class=VALUES(student_class)',student_rows)
        # Reload IDs after upsert.
        students={str(x['roll_no']).casefold():x['student_id'] for x in _select_by_in(cur,'SELECT student_id,roll_no,enrollment_no FROM students WHERE roll_no IN ({ph})',rolls)}
        # Catalogs are inserted once and resolved once.
        subject_names=list(dict.fromkeys(r['subject_name'] for r in clean)); sem_names=list(dict.fromkeys(r['semester_name'] for r in clean)); exam_names=list(dict.fromkeys(r['exam_name'] for r in clean))
        subject_ids=_fetch_map(cur,'subjects','subject_id','subject_name',subject_names)
        sem_ids=_fetch_map(cur,'semesters','semester_id','semester_name',sem_names)
        exam_ids=_fetch_map(cur,'exam_types','exam_id','exam_name',exam_names)
        ss_pairs=[]
        for r in clean:
            sid=students[r['roll_no'].casefold()]; subid=subject_ids[r['subject_name'].casefold()]; ss_pairs.append((sid,subid))
        _executemany_chunks(cur,'INSERT IGNORE INTO student_subjects(student_id,subject_id) VALUES(%s,%s)',list(dict.fromkeys(ss_pairs)))
        sid_list=list(dict.fromkeys(x[0] for x in ss_pairs)); sub_list=list(dict.fromkeys(x[1] for x in ss_pairs))
        ss_map={}
        for sb in _chunks(sid_list):
            ph1=','.join(['%s']*len(sb)); ph2=','.join(['%s']*len(sub_list))
            cur.execute(f'SELECT student_subject_id,student_id,subject_id FROM student_subjects WHERE student_id IN ({ph1}) AND subject_id IN ({ph2})',sb+sub_list)
            ss_map.update({(x['student_id'],x['subject_id']):x['student_subject_id'] for x in cur.fetchall()})
        mark_rows=[]
        for r in clean:
            sid=students[r['roll_no'].casefold()]; subid=subject_ids[r['subject_name'].casefold()]; ssid=ss_map[(sid,subid)]
            mark_rows.append((ssid,sem_ids[r['semester_name'].casefold()],exam_ids[r['exam_name'].casefold()],safe_float(r['obtained_marks']),safe_float(r['maximum_marks'])))
        _executemany_chunks(cur,'INSERT INTO marks(student_subject_id,semester_id,exam_id,obtained_marks,maximum_marks) VALUES(%s,%s,%s,%s,%s) ON DUPLICATE KEY UPDATE obtained_marks=VALUES(obtained_marks),maximum_marks=VALUES(maximum_marks)',mark_rows)
        conn.commit(); return len(mark_rows)
    except Exception:
        if conn:
            try: conn.rollback()
            except Exception: pass
        raise
    finally:
        if cur:
            try: cur.close()
            except Exception: pass
        if conn:
            try: conn.close()
            except Exception: pass

@app.get('/marks/import/data')
@login_required
@role_required('admin','teacher')
def marks_import_data():
    token=session.get('import_token'); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if not token or not os.path.exists(path): return jsonify(ok=False,error='Import session expired.'),404
    try:
        offset=max(0,request.args.get('offset',0,type=int)); limit=min(100, max(1,request.args.get('limit',60,type=int)))
        with open(path,encoding='utf-8') as fh: payload=json.load(fh)
        rows=payload.get('rows',[]); page=[]
        for i,r in enumerate(rows[offset:offset+limit],offset):
            item=dict(r); item['_index']=i; page.append(item)
        return jsonify(ok=True,rows=page,total=len(rows),offset=offset,limit=limit)
    except Exception as exc: return jsonify(ok=False,error=str(exc)),400

@app.route('/marks/import/preview',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def import_preview():
    token=session.get('import_token'); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if not token or not os.path.exists(path):
        flash('Import session expired.','danger'); return redirect(url_for('import_marks'))
    with open(path,encoding='utf-8') as fh: payload=json.load(fh)
    original=payload.get('rows',[])
    if request.method=='GET': return render_template('import_preview.html',rows=[],row_total=len(original),import_info=payload)
    edited=[]
    try:
        edited=json.loads(request.form.get('rows_json','[]')); saved=_save_marks_rows(payload,edited)
        os.remove(path); session.pop('import_token',None)
        flash(f'{saved} mark row(s) imported successfully. Multiple tests in the same file were saved separately.','success')
        return redirect(url_for('results'))
    except Exception as exc:
        app.logger.exception('Import save failed: %s',exc)
        return render_template('import_preview.html',rows=edited or original[:60],row_total=len(original),import_info=payload,save_error=str(exc))

@app.post('/marks/import/save-chunk')
@login_required
@role_required('admin','teacher')
def marks_import_save_chunk():
    token=session.get('import_token'); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if not token or not os.path.exists(path): return jsonify(ok=False,error='Import session expired. Please upload the file again.'),400
    try:
        data=request.get_json(silent=True) or {}; page=int(data.get('page_index',0)); rows=data.get('rows',[]); deleted=data.get('deleted_indices',[]); added=data.get('added_rows',[])
        if page<0 or not isinstance(rows,list) or not isinstance(deleted,list) or not isinstance(added,list): raise ValueError('Invalid import changes.')
        with open(path,encoding='utf-8') as fh: payload=json.load(fh)
        pages=payload.get('edit_pages',{})
        for r in rows:
            if not isinstance(r,dict) or r.get('_index') is None: continue
            rid=int(r.get('_index')); pages.setdefault(str(rid//60),{})[str(rid)]=r
        payload['edit_pages']=pages
        deleted_all=set(payload.get('deleted_indices',[])); deleted_all.update(int(x) for x in deleted if str(x).isdigit()); payload['deleted_indices']=sorted(deleted_all)
        payload['added_rows']=added
        with open(path,'w',encoding='utf-8') as fh: json.dump(payload,fh,ensure_ascii=False,separators=(',',':'))
        return jsonify(ok=True,page=page)
    except Exception as exc: return jsonify(ok=False,error=str(exc)),400

@app.post('/marks/import/finalize')
@login_required
@role_required('admin','teacher')
def marks_import_finalize():
    token=session.get('import_token'); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if not token or not os.path.exists(path): return jsonify(ok=False,error='Import session expired. Please upload the file again.'),400
    try:
        with open(path,encoding='utf-8') as fh: payload=json.load(fh)
        original=payload.get('rows',[]); pages=payload.get('edit_pages',{}); deleted={int(x) for x in payload.get('deleted_indices',[]) if str(x).isdigit()}; edited=[]
        for i,row in enumerate(original):
            if i in deleted: continue
            page=i//60; replacement=pages.get(str(page),{}).get(str(i),row)
            if replacement and not replacement.get('_deleted'): edited.append({k:v for k,v in replacement.items() if k!='_index'})
        edited.extend(payload.get('added_rows',[]))
        saved=_save_marks_rows(payload,edited)
        os.remove(path); session.pop('import_token',None)
        flash(f'{saved} mark row(s) imported successfully. Multiple tests in the same file were saved separately.','success')
        return jsonify(ok=True,saved=saved,redirect=url_for('results'))
    except Exception as exc:
        app.logger.exception('Chunked import finalize failed: %s',exc); return jsonify(ok=False,error=str(exc)),400

@app.post('/marks/import/cancel-chunked')
@login_required
def cancel_chunked_import():
    token=session.pop('import_token',None); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if token and os.path.exists(path): os.remove(path)
    return jsonify(ok=True)

@app.post('/marks/import/cancel')
@login_required
def cancel_import():
    token=session.pop('import_token',None); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if token and os.path.exists(path): os.remove(path)
    return redirect(url_for('import_marks'))

@app.route('/marks/import/template/<fmt>')
@login_required
@role_required('admin','teacher')
def import_template(fmt):
    headers=['Roll No','Enrollment No','Student Name','Class','Subject','Semester','Exam Type','Obtained Marks','Maximum Marks']; example=['1','1','Student A','FY','Mathematics','Semester 1','Unit Test 1','18','20']
    if fmt=='csv':
        s=io.StringIO(newline=''); w=csv.writer(s); w.writerow(headers); w.writerow(example); return send_file(io.BytesIO(s.getvalue().encode('utf-8-sig')),as_attachment=True,download_name='marks_import_template.csv',mimetype='text/csv')
    if fmt=='xlsx':
        wb=Workbook(); ws=wb.active; ws.append(headers); ws.append(example); bio=io.BytesIO(); wb.save(bio); bio.seek(0); return send_file(bio,as_attachment=True,download_name='marks_import_template.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    return redirect(url_for('import_marks'))

# ---------------- RESULTS ----------------
def order_students(rows):
    return sorted(rows,key=lambda r:(0,int(r['roll_no'])) if str(r.get('roll_no','')).isdigit() else (1,str(r.get('roll_no','')).casefold()))

def mark_rows_for(scope,sp,sem=None,exam=None):
    cond=[scope]; params=list(sp)
    if sem: cond.append('m.semester_id=%s'); params.append(sem)
    if exam: cond.append('m.exam_id=%s'); params.append(exam)
    return query_db(f'''SELECT s.student_id,s.roll_no,s.enrollment_no,s.student_name,s.student_class,sub.subject_id,sub.subject_name,m.obtained_marks,m.maximum_marks,m.semester_id,m.exam_id,e.exam_name,sem.semester_name FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id JOIN students s ON s.student_id=ss.student_id JOIN subjects sub ON sub.subject_id=ss.subject_id JOIN semesters sem ON sem.semester_id=m.semester_id JOIN exam_types e ON e.exam_id=m.exam_id WHERE {' AND '.join(cond)} ORDER BY s.student_class,s.student_name,sub.subject_name''',params)

def grouped_matrix(data, single_exam=False):
    if single_exam:
        columns=sorted({(x['subject_id'],x['subject_name']) for x in data},key=lambda x:x[1].casefold())
        def key(r): return r['subject_id']
    else:
        columns=sorted({(x['exam_id'],x['subject_id'],x['exam_name'],x['subject_name']) for x in data},key=lambda x:(x[0],x[3].casefold()))
        def key(r): return (r['exam_id'],r['subject_id'])
    grouped={}
    for r in data:
        g=grouped.setdefault(r['student_id'],{'student_id':r['student_id'],'roll_no':r['roll_no'],'enrollment_no':r['enrollment_no'],'student_name':r['student_name'],'student_class':r['student_class'],'marks':{},'obtained':0.0,'maximum':0.0})
        g['marks'][key(r)]=r; g['obtained']+=float(r['obtained_marks']); g['maximum']+=float(r['maximum_marks'])
    rows=order_students(list(grouped.values()))
    for r in rows: r['percentage']=round(r['obtained']/r['maximum']*100,2) if r['maximum'] else 0; r['grade']=grade_for(r['percentage'])
    return rows,columns

def grade_for(p): return 'A+' if p>=90 else 'A' if p>=80 else 'B+' if p>=70 else 'B' if p>=60 else 'C' if p>=50 else 'D' if p>=40 else 'F'


def performance_remark(percentage, attendance_pct=None):
    """Teacher-friendly automatic observation; never replaces a teacher's own remark."""
    try: p=float(percentage or 0)
    except Exception: p=0.0
    if p >= 90: remark = "Excellent academic performance. The student is demonstrating strong understanding and consistency."
    elif p >= 80: remark = "Very good performance. The student is progressing well; maintain consistency and strengthen weaker areas."
    elif p >= 70: remark = "Good performance. With focused practice in weaker subjects, the student can improve further."
    elif p >= 60: remark = "Satisfactory performance. More regular revision and practice are recommended."
    elif p >= 50: remark = "Needs improvement. The student should receive additional practice, revision and teacher support."
    else: remark = "Academic performance needs immediate attention. A focused improvement plan and regular follow-up are recommended."
    if attendance_pct is not None:
        try:
            a=float(attendance_pct)
            if a < 75: remark += " Attendance is below the detention threshold and requires attention."
            elif a < 85: remark += " Attendance is acceptable but should be improved for better continuity."
            else: remark += " Attendance is supporting regular academic participation."
        except Exception:
            pass
    return remark

@app.route('/results')
@login_required
def results():
    u=current_user(); scope,sp=student_scope_condition(); semesters=semester_choices(); exams=exam_choices(); classes=allowed_classes(); sid=request.args.get('student_id',type=int); sem=request.args.get('semester_id',type=int); exam=request.args.get('exam_id',type=int); cls=request.args.get('student_class','').strip(); search=request.args.get('search','').strip()
    if u['role']=='student': sid=u['student_id']; cls=u['student_class']
    if u['role']=='teacher' and (not cls or cls not in classes): cls=active_class()
    if cls and not enforce_class(cls): cls=''
    # A selected class is authoritative: a student from another class is never shown.
    if sid:
        sr=query_db('SELECT student_class FROM students WHERE student_id=%s',(sid,),fetchone=True)
        if not sr or not enforce_class(sr['student_class']) or (cls and sr['student_class']!=cls): sid=None
    cond=scope; params=list(sp)
    if sid: cond+=' AND s.student_id=%s'; params.append(sid)
    if cls: cond+=' AND s.student_class=%s'; params.append(cls)
    if search: cond+=' AND (s.roll_no LIKE %s OR s.enrollment_no LIKE %s)'; params += [f'%{search}%',f'%{search}%']
    data=mark_rows_for(cond,params,sem,exam); rows,subjects=grouped_matrix(data, single_exam=bool(exam))
    opt_cond=scope; opt_params=list(sp)
    if cls: opt_cond+=' AND s.student_class=%s'; opt_params.append(cls)
    if search: opt_cond+=' AND (s.roll_no LIKE %s OR s.enrollment_no LIKE %s OR s.student_name LIKE %s)'; opt_params += [f'%{search}%',f'%{search}%',f'%{search}%']
    student_options=query_db(f'SELECT s.student_id,s.roll_no,s.enrollment_no,s.student_name,s.student_class FROM students s WHERE {opt_cond} ORDER BY CASE WHEN s.roll_no REGEXP "^[0-9]+$" THEN 0 ELSE 1 END, CASE WHEN s.roll_no REGEXP "^[0-9]+$" THEN CAST(s.roll_no AS UNSIGNED) ELSE 0 END,s.roll_no,s.student_name',opt_params)
    # Class-level dashboard summary.
    class_avg=round(sum(float(r['percentage']) for r in rows)/len(rows),2) if rows else 0
    grade_counts={}
    for r in rows: grade_counts[r['grade']]=grade_counts.get(r['grade'],0)+1
    subject_avgs=[]
    for col in subjects:
        vals=[]
        for r in rows:
            mk=r['marks'].get(col if len(col)==2 else (col[0],col[1]))
            if mk and float(mk['maximum_marks']): vals.append(float(mk['obtained_marks'])/float(mk['maximum_marks'])*100)
        if vals: subject_avgs.append({'label':col[1] if len(col)==2 else col[3],'value':round(sum(vals)/len(vals),2)})
    return render_template('results.html',rows=rows,subjects=subjects,students=student_options,semesters=semesters,exams=exams,classes=classes,selected_student=sid,selected_semester=sem,selected_exam=exam,selected_class=cls,search=search,class_avg=class_avg,grade_counts=grade_counts,subject_avgs=subject_avgs)

@app.post('/students/<int:student_id>/remark')
@login_required
@role_required('admin','teacher')
def save_student_remark(student_id):
    student=query_db('SELECT student_class FROM students WHERE student_id=%s',(student_id,),fetchone=True)
    if not student or not enforce_class(student['student_class']):
        flash('Student is outside your access.','danger')
        return redirect(url_for('analysis'))
    remark=request.form.get('teacher_remark','').strip()[:500]
    query_db('UPDATE students SET teacher_remark=%s WHERE student_id=%s',(remark,student_id),commit=True)
    flash('Teacher remark saved.','success')
    return redirect(url_for('analysis'))

# ---------------- ATTENDANCE ----------------
def month_id(conn,cur,key,label=None):
    cur.execute('SELECT attendance_month_id FROM attendance_months WHERE month_key=%s',(key,)); x=cur.fetchone()
    if x: return x['attendance_month_id']
    cur.execute('INSERT INTO attendance_months(month_key,month_label) VALUES(%s,%s)',(key,label or key)); return cur.lastrowid

@app.route('/attendance/import',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def import_attendance():
    if request.method=='POST':
        f=request.files.get('attendance_file'); month=request.form.get('month','').strip(); use_file_month=request.form.get('use_file_month')=='1'
        if not f or not f.filename: flash('Select an attendance file.','danger'); return redirect(url_for('import_attendance'))
        if not month and not use_file_month: flash('Select the attendance month or enable Month from file.','danger'); return redirect(url_for('import_attendance'))
        try:
            rows=normalize(read_file(f),ATT_MAP)
            validate_import_row_count(len(rows))
            if not rows: raise ValueError('No readable attendance rows found.')
            token=uuid.uuid4().hex; payload={'kind':'attendance','month':month,'use_file_month':use_file_month,'rows':rows}
            with open(os.path.join(IMPORT_DIR,token+'.json'),'w',encoding='utf-8') as out: json.dump(payload,out,ensure_ascii=False)
            session['attendance_import_token']=token; return redirect(url_for('attendance_preview'))
        except Exception as exc: flash(f'Attendance file could not be read: {exc}','danger')
    return render_template('attendance_import.html')

def _validate_attendance_rows(edited,payload):
    fields=['roll_no','enrollment_no','student_name','student_class','subject_name','month','working_days','present_days','attendance_pct','notes']
    errors=[]; clean=[]
    for i,item in enumerate(edited,1):
        r={k:str(item.get(k,'') or '').strip() for k in fields}
        if not any(r.values()): continue
        r['enrollment_no']=r['enrollment_no'] or r['roll_no']; r['month']=r['month'] or payload['month']
        if not r['roll_no'] or not r['student_name']: errors.append(f'Row {i}: Roll No and Student Name are required.'); continue
        wd=safe_float(r['working_days']); pd=safe_float(r['present_days']); pct=safe_float(r['attendance_pct'])
        if wd is None or wd<=0:
            if pct is not None and pd is not None and pd>=0: wd=100; pd=wd*pct/100
            else: errors.append(f'Row {i}: enter Working Days and Present Days.'); continue
        if pd is None and pct is not None: pd=wd*pct/100
        if pd is None or pd<0 or pd>wd: errors.append(f'Row {i}: Present Days must be between 0 and Working Days.'); continue
        if not enforce_class(r['student_class'] or 'Not Provided') and current_user()['role']!='admin': errors.append(f'Row {i}: class is outside your assigned classes.'); continue
        r['working_days']=str(wd); r['present_days']=str(pd); r['attendance_pct']=str(round(pd/wd*100,2)); clean.append(r)
    return clean,errors

def _save_attendance_rows(payload,edited):
    clean,errors=_validate_attendance_rows(edited,payload)
    if errors: raise ValueError(' | '.join(errors[:12]))
    if not clean: return 0
    conn=cur=None
    try:
        conn=get_db(); conn.start_transaction(); cur=conn.cursor(dictionary=True)
        role=current_user()['role']
        rolls=list(dict.fromkeys(r['roll_no'] for r in clean)); ens=list(dict.fromkeys(r['enrollment_no'] or r['roll_no'] for r in clean))
        by_roll={str(x['roll_no']).casefold():x for x in _select_by_in(cur,'SELECT * FROM students WHERE roll_no IN ({ph})',rolls)}
        by_en={str(x['enrollment_no']).casefold():x for x in _select_by_in(cur,'SELECT * FROM students WHERE enrollment_no IN ({ph})',ens) if x.get('enrollment_no')}
        student_rows=[]
        for r in clean:
            a=by_roll.get(r['roll_no'].casefold()); b=by_en.get((r['enrollment_no'] or r['roll_no']).casefold())
            if a and b and a['student_id']!=b['student_id']: raise ValueError(f"Roll {r['roll_no']} and Enrollment {r['enrollment_no']} identify different students.")
            stu=a or b; cls=r['student_class'] or (stu['student_class'] if stu else 'Not Provided')
            if not enforce_class(cls) and role!='admin': raise ValueError(f'Class {cls} is outside your assigned classes.')
            student_rows.append((r['roll_no'],r['enrollment_no'] or r['roll_no'],r['student_name'],cls))
        _executemany_chunks(cur,'INSERT INTO students(roll_no,enrollment_no,student_name,student_class) VALUES(%s,%s,%s,%s) ON DUPLICATE KEY UPDATE enrollment_no=VALUES(enrollment_no),student_name=VALUES(student_name),student_class=VALUES(student_class)',student_rows)
        students={str(x['roll_no']).casefold():x['student_id'] for x in _select_by_in(cur,'SELECT student_id,roll_no FROM students WHERE roll_no IN ({ph})',rolls)}
        months=[]; subjects=[]
        for r in clean:
            month=r['month'] or payload['month']
            try: key=datetime.strptime(month,'%Y-%m').strftime('%Y-%m'); label=datetime.strptime(month,'%Y-%m').strftime('%B %Y')
            except ValueError:
                try: d=datetime.strptime(month,'%B %Y'); key=d.strftime('%Y-%m'); label=d.strftime('%B %Y')
                except ValueError: key='LBL-'+re.sub(r'[^A-Za-z0-9]','',month)[:6] + '-' + str(abs(hash(month))%100000); label=month
            months.append((key,label))
            if r['subject_name']: subjects.append(r['subject_name'])
        _executemany_chunks(cur,'INSERT IGNORE INTO attendance_months(month_key,month_label) VALUES(%s,%s)',list(dict.fromkeys(months)))
        mids={}
        for mb in _chunks([x[0] for x in dict.fromkeys(months)]):
            ph=','.join(['%s']*len(mb)); cur.execute(f'SELECT attendance_month_id,month_key FROM attendance_months WHERE month_key IN ({ph})',mb); mids.update({x['month_key']:x['attendance_month_id'] for x in cur.fetchall()})
        subject_ids=_fetch_map(cur,'subjects','subject_id','subject_name',subjects)
        # Ensure subject assignments exist for subject-specific attendance.
        ss_pairs=[]
        attendance_rows=[]
        for r,(key,label) in zip(clean,months):
            sid=students[r['roll_no'].casefold()]; subid=subject_ids.get(r['subject_name'].casefold()) if r['subject_name'] else None
            if subid is not None: ss_pairs.append((sid,subid))
            attendance_rows.append((sid,mids[key],subid,safe_float(r['working_days']),safe_float(r['present_days']),r['notes']))
        if ss_pairs: _executemany_chunks(cur,'INSERT IGNORE INTO student_subjects(student_id,subject_id) VALUES(%s,%s)',list(dict.fromkeys(ss_pairs)))
        _executemany_chunks(cur,'INSERT INTO attendance(student_id,attendance_month_id,subject_id,working_days,present_days,notes) VALUES(%s,%s,%s,%s,%s,%s) ON DUPLICATE KEY UPDATE working_days=VALUES(working_days),present_days=VALUES(present_days),notes=VALUES(notes)',attendance_rows)
        conn.commit(); return len(attendance_rows)
    except Exception:
        if conn:
            try: conn.rollback()
            except Exception: pass
        raise
    finally:
        if cur:
            try: cur.close()
            except Exception: pass
        if conn:
            try: conn.close()
            except Exception: pass

@app.get('/attendance/import/data')
@login_required
@role_required('admin','teacher')
def attendance_import_data():
    token=session.get('attendance_import_token'); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if not token or not os.path.exists(path): return jsonify(ok=False,error='Attendance import session expired.'),404
    try:
        offset=max(0,request.args.get('offset',0,type=int)); limit=min(100,max(1,request.args.get('limit',60,type=int)))
        with open(path,encoding='utf-8') as fh: payload=json.load(fh)
        rows=payload.get('rows',[]); page=[]
        for i,r in enumerate(rows[offset:offset+limit],offset): item=dict(r); item['_index']=i; page.append(item)
        return jsonify(ok=True,rows=page,total=len(rows),offset=offset,limit=limit)
    except Exception as exc: return jsonify(ok=False,error=str(exc)),400

@app.route('/attendance/import/preview',methods=['GET','POST'])
@login_required
@role_required('admin','teacher')
def attendance_preview():
    token=session.get('attendance_import_token'); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if not token or not os.path.exists(path): flash('Attendance import session expired.','danger'); return redirect(url_for('import_attendance'))
    with open(path,encoding='utf-8') as fh: payload=json.load(fh)
    original=payload.get('rows',[])
    if request.method=='GET': return render_template('attendance_preview.html',rows=[],row_total=len(original),import_info=payload)
    edited=[]
    try:
        edited=json.loads(request.form.get('rows_json','[]')); saved=_save_attendance_rows(payload,edited); os.remove(path); session.pop('attendance_import_token',None); flash(f'{saved} attendance row(s) saved.','success'); return redirect(url_for('attendance_dashboard'))
    except Exception as exc:
        return render_template('attendance_preview.html',rows=edited or original[:60],row_total=len(original),import_info=payload,save_error=str(exc))

@app.post('/attendance/import/save-chunk')
@login_required
@role_required('admin','teacher')
def attendance_import_save_chunk():
    token=session.get('attendance_import_token'); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if not token or not os.path.exists(path): return jsonify(ok=False,error='Attendance import session expired. Please upload the file again.'),400
    try:
        data=request.get_json(silent=True) or {}; page=int(data.get('page_index',0)); rows=data.get('rows',[]); deleted=data.get('deleted_indices',[]); added=data.get('added_rows',[])
        if page<0 or not isinstance(rows,list) or not isinstance(deleted,list) or not isinstance(added,list): raise ValueError('Invalid attendance changes.')
        with open(path,encoding='utf-8') as fh: payload=json.load(fh)
        pages=payload.get('edit_pages',{})
        for r in rows:
            if not isinstance(r,dict) or r.get('_index') is None: continue
            rid=int(r.get('_index')); pages.setdefault(str(rid//60),{})[str(rid)]=r
        payload['edit_pages']=pages
        deleted_all=set(payload.get('deleted_indices',[])); deleted_all.update(int(x) for x in deleted if str(x).isdigit()); payload['deleted_indices']=sorted(deleted_all); payload['added_rows']=added
        with open(path,'w',encoding='utf-8') as fh: json.dump(payload,fh,ensure_ascii=False,separators=(',',':'))
        return jsonify(ok=True,page=page)
    except Exception as exc: return jsonify(ok=False,error=str(exc)),400

@app.post('/attendance/import/finalize')
@login_required
@role_required('admin','teacher')
def attendance_import_finalize():
    token=session.get('attendance_import_token'); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if not token or not os.path.exists(path): return jsonify(ok=False,error='Attendance import session expired. Please upload the file again.'),400
    try:
        with open(path,encoding='utf-8') as fh: payload=json.load(fh)
        original=payload.get('rows',[]); pages=payload.get('edit_pages',{}); deleted={int(x) for x in payload.get('deleted_indices',[]) if str(x).isdigit()}; edited=[]
        for i,row in enumerate(original):
            if i in deleted: continue
            page=i//60; replacement=pages.get(str(page),{}).get(str(i),row)
            if replacement and not replacement.get('_deleted'): edited.append({k:v for k,v in replacement.items() if k!='_index'})
        edited.extend(payload.get('added_rows',[]))
        saved=_save_attendance_rows(payload,edited); os.remove(path); session.pop('attendance_import_token',None); flash(f'{saved} attendance row(s) saved.','success'); return jsonify(ok=True,saved=saved,redirect=url_for('attendance_dashboard'))
    except Exception as exc:
        app.logger.exception('Chunked attendance finalize failed: %s',exc); return jsonify(ok=False,error=str(exc)),400

@app.post('/attendance/import/cancel-chunked')
@login_required
def cancel_attendance_chunked():
    token=session.pop('attendance_import_token',None); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if token and os.path.exists(path): os.remove(path)
    return jsonify(ok=True)

@app.post('/attendance/import/cancel')
@login_required
def cancel_attendance_import():
    token=session.pop('attendance_import_token',None); path=os.path.join(IMPORT_DIR,(token or '')+'.json')
    if token and os.path.exists(path): os.remove(path)
    return redirect(url_for('import_attendance'))

# ---------------- ATTENDANCE DASHBOARD / STUDENT ANALYSIS ----------------
def attendance_summary(student_id=None,class_name=None):
    cond=[]; params=[]; scope,sp=student_scope_condition(); cond.append(scope); params += sp
    if student_id: cond.append('s.student_id=%s'); params.append(student_id)
    if class_name: cond.append('s.student_class=%s'); params.append(class_name)
    where=' AND '.join(cond)
    overall=query_db(f'''SELECT ROUND(SUM(a.present_days)/NULLIF(SUM(a.working_days),0)*100,2) pct,SUM(a.present_days) present,SUM(a.working_days) working FROM attendance a JOIN students s ON s.student_id=a.student_id WHERE {where}''',params,fetchone=True)
    monthly=query_db(f'''SELECT am.month_key,am.month_label,ROUND(SUM(a.present_days)/NULLIF(SUM(a.working_days),0)*100,2) pct FROM attendance a JOIN attendance_months am ON am.attendance_month_id=a.attendance_month_id JOIN students s ON s.student_id=a.student_id WHERE {where} GROUP BY am.attendance_month_id,am.month_key,am.month_label ORDER BY am.month_key''',params)
    subject=query_db(f'''SELECT COALESCE(sub.subject_name,'Overall') subject_name,ROUND(SUM(a.present_days)/NULLIF(SUM(a.working_days),0)*100,2) pct FROM attendance a JOIN students s ON s.student_id=a.student_id LEFT JOIN subjects sub ON sub.subject_id=a.subject_id WHERE {where} GROUP BY sub.subject_id,sub.subject_name ORDER BY sub.subject_name''',params)
    return overall or {'pct':0,'present':0,'working':0},monthly,subject

@app.route('/attendance')
@login_required
def attendance_dashboard():
    u=current_user(); classes=allowed_classes(); cls=request.args.get('student_class','').strip(); sid=request.args.get('student_id',type=int); search=request.args.get('search','').strip()
    if u['role']=='student': sid=u['student_id']; cls=u['student_class']
    if u['role']=='teacher' and (not cls or cls not in classes): cls=active_class()
    if cls and not enforce_class(cls): cls=''
    if search and not sid:
        found=query_db('SELECT student_id,student_class FROM students WHERE (roll_no=%s OR enrollment_no=%s) LIMIT 1',(search,search),fetchone=True)
        if found and enforce_class(found['student_class']): sid=found['student_id']
    overall,monthly,subject=attendance_summary(sid,cls)
    students=query_db('SELECT student_id,roll_no,enrollment_no,student_name,student_class FROM students WHERE student_class IN ('+','.join(['%s']*len(classes))+') ORDER BY student_class,student_name',classes) if classes else []
    threshold=institution_config()['detention_threshold']; status='At Risk of Detention' if float(overall.get('pct') or 0)<threshold else 'Above Detention Threshold'
    detention=[]
    if cls and not sid:
        for st in students:
            if st['student_class']!=cls: continue
            ao,_,_=attendance_summary(st['student_id'],None); pct=float(ao.get('pct') or 0)
            if ao.get('working') and pct<threshold: detention.append({'student':st,'pct':pct})
        detention=sorted(detention,key=lambda x:x['pct'])
    return render_template('attendance_dashboard.html',overall=overall,monthly=monthly,subject=subject,classes=classes,students=students,selected_class=cls,selected_student=sid,search=search,threshold=threshold,status=status,detention=detention)

# ---------------- ANALYSIS ----------------
def student_analysis_data(sid,sem=None,exam=None):
    cond='s.student_id=%s'; params=[sid]
    if sem: cond+=' AND m.semester_id=%s'; params.append(sem)
    if exam: cond+=' AND m.exam_id=%s'; params.append(exam)
    rows=query_db(f'''SELECT sub.subject_name,m.obtained_marks obtained,m.maximum_marks maximum,ROUND(m.obtained_marks/m.maximum_marks*100,2) percentage,e.exam_name,sem.semester_name FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id JOIN students s ON s.student_id=ss.student_id JOIN subjects sub ON sub.subject_id=ss.subject_id JOIN semesters sem ON sem.semester_id=m.semester_id JOIN exam_types e ON e.exam_id=m.exam_id WHERE {cond} ORDER BY e.exam_id,sem.semester_id,sub.subject_name''',params)
    return rows

@app.route('/analysis',methods=['GET','POST'])
@login_required
def analysis():
    u=current_user(); semesters=semester_choices(); exams=exam_choices(); classes=allowed_classes(); sid=request.values.get('student_id',type=int); sem=request.values.get('semester_id',type=int); exam=request.values.get('exam_id',type=int); cls=request.values.get('student_class','').strip(); search=request.values.get('search','').strip(); view=request.values.get('view','overall_performance')
    if u['role']=='student': sid=u['student_id']; cls=u['student_class']
    if u['role']=='teacher' and (not cls or cls not in classes): cls=active_class()
    if search and not sid:
        found=query_db('SELECT student_id,student_class FROM students WHERE roll_no=%s OR enrollment_no=%s LIMIT 1',(search,search),fetchone=True)
        if found and enforce_class(found['student_class']) and (not cls or found['student_class']==cls): sid=found['student_id']
    if sid:
        st=query_db('SELECT * FROM students WHERE student_id=%s',(sid,),fetchone=True)
        if not st or not enforce_class(st['student_class']) or (cls and st['student_class']!=cls): sid=None; st=None
    else: st=None
    student_data=student_analysis_data(sid,sem,exam) if sid else []
    class_grouped=[]; class_subjects=[]
    if cls and not sid:
        cond='s.student_class=%s'; cp=[cls]
        if sem: cond+=' AND m.semester_id=%s'; cp.append(sem)
        if exam: cond+=' AND m.exam_id=%s'; cp.append(exam)
        cdata=query_db(f'''SELECT s.student_id,s.roll_no,s.enrollment_no,s.student_name,s.student_class,sub.subject_id,sub.subject_name,m.obtained_marks,m.maximum_marks,m.exam_id,e.exam_name,m.semester_id,sem.semester_name FROM marks m JOIN student_subjects ss ON ss.student_subject_id=m.student_subject_id JOIN students s ON s.student_id=ss.student_id JOIN subjects sub ON sub.subject_id=ss.subject_id JOIN exam_types e ON e.exam_id=m.exam_id JOIN semesters sem ON sem.semester_id=m.semester_id WHERE {cond} ORDER BY s.student_name,sub.subject_name,e.exam_id''',cp)
        class_subjects=sorted({(r['subject_id'],r['subject_name']) for r in cdata},key=lambda x:x[1].casefold())
        by={}
        for r in cdata:
            g=by.setdefault(r['student_id'],{'student_id':r['student_id'],'roll_no':r['roll_no'],'enrollment_no':r['enrollment_no'],'student_name':r['student_name'],'student_class':r['student_class'],'marks':{},'obtained':0,'maximum':0})
            g['marks'][r['subject_id']]=r; g['obtained']+=float(r['obtained_marks']); g['maximum']+=float(r['maximum_marks'])
        class_grouped=order_students(list(by.values()))
        for g in class_grouped: g['percentage']=round(g['obtained']/g['maximum']*100,2) if g['maximum'] else 0; g['grade']=grade_for(g['percentage'])
    ob=sum(float(x['obtained']) for x in student_data); mx=sum(float(x['maximum']) for x in student_data); pct=round(ob/mx*100,2) if mx else 0
    attendance={'pct':0,'present':0,'working':0}; monthly=[]; subject_att=[]
    if sid and view=='overall_performance': attendance,monthly,subject_att=attendance_summary(sid,None)
    # Per-test summaries make Unit 1 -> Unit 2 -> Semester easy to read.
    test_groups=[]
    for r in student_data:
        key=(r['exam_name'],r['semester_name']); g=next((x for x in test_groups if x['key']==key),None)
        if not g: g={'key':key,'exam_name':r['exam_name'],'semester_name':r['semester_name'],'rows':[],'obtained':0,'maximum':0}
        g['rows'].append(r); g['obtained']+=float(r['obtained']); g['maximum']+=float(r['maximum'])
        if g not in test_groups: test_groups.append(g)
    for g in test_groups: g['percentage']=round(g['obtained']/g['maximum']*100,2) if g['maximum'] else 0; g['grade']=grade_for(g['percentage'])
    trend=[{'label':g['exam_name'],'value':g['percentage']} for g in test_groups]
    class_overall={'pct':0,'present':0,'working':0}; class_monthly=[]; class_subject_att=[]
    class_subject_avgs=[]; class_grade_counts={}; class_ob=class_mx=0
    if cls and not sid:
        class_overall,class_monthly,class_subject_att=attendance_summary(None,cls)
        for g in class_grouped:
            class_ob += float(g.get('obtained') or 0); class_mx += float(g.get('maximum') or 0)
            class_grade_counts[g['grade']]=class_grade_counts.get(g['grade'],0)+1
        for subject_id,subject_name in class_subjects:
            vals=[]
            for g in class_grouped:
                mk=g['marks'].get(subject_id)
                if mk and float(mk.get('maximum_marks') or 0): vals.append(float(mk['obtained_marks'])/float(mk['maximum_marks'])*100)
            if vals: class_subject_avgs.append({'label':subject_name,'value':round(sum(vals)/len(vals),2)})
    class_percentage=round(class_ob/class_mx*100,2) if class_mx else 0
    prediction=None
    if len(trend)>=2:
        diffs=[trend[i]['value']-trend[i-1]['value'] for i in range(1,len(trend))]; prediction=round(max(0,min(100,trend[-1]['value']+statistics.mean(diffs))),2)
    opt_cond='student_class IN ('+','.join(['%s']*len(classes))+')'; opt_params=list(classes)
    if cls: opt_cond+=' AND student_class=%s'; opt_params.append(cls)
    if search: opt_cond+=' AND (roll_no LIKE %s OR enrollment_no LIKE %s OR student_name LIKE %s)'; opt_params += [f'%{search}%',f'%{search}%',f'%{search}%']
    analysis_students=query_db(f'SELECT student_id,roll_no,enrollment_no,student_name,student_class FROM students WHERE {opt_cond} ORDER BY CASE WHEN roll_no REGEXP "^[0-9]+$" THEN 0 ELSE 1 END, CASE WHEN roll_no REGEXP "^[0-9]+$" THEN CAST(roll_no AS UNSIGNED) ELSE 0 END,roll_no,student_name',opt_params) if classes else []
    auto_remark=performance_remark(pct, attendance.get('pct') if sid and view=='overall_performance' else None)
    return render_template('analysis.html',students=analysis_students,semesters=semesters,exams=exams,classes=classes,student=st,selected_student=sid,selected_semester=sem,selected_exam=exam,selected_class=cls,search=search,view=view,data=student_data,total_obtained=ob,total_maximum=mx,percentage=pct,grade=grade_for(pct),test_groups=test_groups,trend=trend,prediction=prediction,attendance=attendance,monthly=monthly,subject_att=subject_att,threshold=institution_config()['detention_threshold'],class_grouped=class_grouped,class_subjects=class_subjects,class_overall=class_overall,class_monthly=class_monthly,class_subject_att=class_subject_att,class_subject_avgs=class_subject_avgs,class_percentage=class_percentage,class_grade_counts=class_grade_counts,auto_remark=auto_remark,teacher_remark=(st.get('teacher_remark') if st else ''))

@app.errorhandler(413)
def request_too_large(_):
    flash('The upload request is too large for the server. Please retry; import review saving is chunked so large marks and attendance datasets can be saved safely.', 'danger')
    return redirect(url_for('index'))

@app.errorhandler(CSRFError)
def csrf_error(_):
    return render_template('error.html',code=400,message='Your security token expired or was missing. Refresh the page and try again.'),400

@app.errorhandler(400)
def bad_request(_):
    return render_template('error.html',code=400,message='The request could not be processed. Please check the submitted data and try again.'),400

@app.errorhandler(401)
def unauthorized(_):
    return render_template('error.html',code=401,message='Please sign in to continue.'),401

@app.errorhandler(403)
def forbidden(_):
    return render_template('error.html',code=403,message='You do not have permission to perform this action.'),403

@app.errorhandler(405)
def method_not_allowed(_):
    return render_template('error.html',code=405,message='That action is not available for this request.'),405

@app.errorhandler(500)
def server_error(_):
    app.logger.exception('Unhandled application error')
    return render_template('error.html',code=500,message='Something went wrong on the server. Please try again.'),500

@app.errorhandler(429)
def too_many_requests(_):
    return render_template('error.html',code=429,message='Too many requests. Please wait a moment and try again.'),429

@app.errorhandler(404)
def not_found(_): flash('The requested page was not found.','danger'); return redirect(url_for('index'))

if __name__=='__main__':
    debug=os.getenv('FLASK_DEBUG','0').lower() in {'1','true','yes'}
    app.run(host=os.getenv('FLASK_HOST','127.0.0.1'),port=int(os.getenv('FLASK_PORT','5000')),debug=debug)
